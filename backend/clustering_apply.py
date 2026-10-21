"""Pembaca & penerap format clustering LOKA-Kin (dipakai CLI dan API).

Modul ini memisahkan logika baca-berkas + terapkan peran dari ``scripts/clustering.py``
agar bisa dipakai dua jalur:

* ``scripts/clustering.py`` — dijalankan di mesin lokal saat server berhenti.
* ``POST /api/admin/clustering/apply`` — unggah berkas dari aplikasi oleh Super Admin.

Semua fungsi bekerja pada daftar dokumen staf (list of dict) sehingga pemanggil
yang menentukan cara menyimpannya (snapshot JSON atau MongoDB).
"""

from __future__ import annotations

import csv
import io
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple, Union

from auth_utils import hash_credential
from clustering_export import (
    DEFAULT_PIN,
    ROLES,
    ROLE_CODE_BY_LABEL,
    ROLE_LABEL_BY_CODE,
    SEPARATE_SUPER_ADMIN_DEFAULT,
    SEPARATE_SUPER_ADMIN_ID,
    SUPERVISOR_ROLES,
    normalize_digits,
    separate_super_admin,
)

TRUE_MARKS = {"1", "01", "x", "X", "✓", "v", "V", "ya", "YA", "yes", "TRUE", "True", "true"}
MAX_UPLOAD_BYTES = 8 * 1024 * 1024  # 8 MB


class TemplateError(ValueError):
    """Berkas unggahan tidak sesuai format clustering."""


# --------------------------------------------------------------------------- #
# Utilitas
# --------------------------------------------------------------------------- #
def id_type_for(nip: str) -> str:
    return {16: "NIK", 18: "NIP"}.get(len(nip), "")


def uniform_pin_fields(pin: str, source: str = "uniform-clustering") -> Dict[str, Any]:
    return {
        "pin_hash": hash_credential(pin),
        "is_activated": True,
        "pin_change_required": False,
        "activation_code_hash": "",
        "activation_code_expires_at": "",
        "activation_code_attempts": 0,
        "activation_code_locked_until": "",
        "login_failed_attempts": 0,
        "login_locked_until": "",
        "pin_source": source,
        "pin_updated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
    }


def is_marked(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return int(value) == 1
    return str(value).strip() in TRUE_MARKS


# --------------------------------------------------------------------------- #
# Membaca berkas
# --------------------------------------------------------------------------- #
def _load_workbook(source: Union[str, Path, bytes, io.BytesIO]):
    from openpyxl import load_workbook

    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    try:
        return load_workbook(source, data_only=True)
    except Exception as exc:  # openpyxl melempar beragam tipe galat
        raise TemplateError(
            "Berkas Excel tidak dapat dibaca. Pastikan berkas .xlsx hasil unduhan "
            "menu Clustering akun, bukan format lain atau berkas yang rusak."
        ) from exc


def parse_workbook(source: Union[str, Path, bytes, io.BytesIO]) -> Tuple[List[Dict[str, Any]], Optional[Dict[str, Any]]]:
    """Baca sheet 'Isi Peran' menjadi daftar peran per akun + blok Super Admin."""
    wb = _load_workbook(source)
    if "Isi Peran" not in wb.sheetnames:
        raise TemplateError(
            f"Sheet 'Isi Peran' tidak ditemukan. Sheet yang ada: {', '.join(wb.sheetnames)}."
        )

    ws = wb["Isi Peran"]
    header_row = None
    for row in range(1, 15):
        if str(ws.cell(row=row, column=2).value or "").strip() == "ID Akun":
            header_row = row
            break
    if header_row is None:
        raise TemplateError("Baris header ('ID Akun') tidak ditemukan pada sheet 'Isi Peran'.")

    role_cols: Dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        label = str(ws.cell(row=header_row, column=col).value or "").strip()
        if label in ROLE_CODE_BY_LABEL:
            role_cols[ROLE_CODE_BY_LABEL[label]] = col
    missing = [label for code, label in ROLES if code not in role_cols]
    if missing:
        raise TemplateError(f"Kolom peran tidak lengkap: {', '.join(missing)}.")

    entries: List[Dict[str, Any]] = []
    blank_streak = 0
    for row in range(header_row + 1, ws.max_row + 1):
        marker = str(ws.cell(row=row, column=1).value or "")
        if marker.startswith("AKUN SUPER ADMIN TERPISAH"):
            break
        ident = str(ws.cell(row=row, column=2).value or "").strip()
        if not ident:
            blank_streak += 1
            if blank_streak >= 5:
                break
            continue
        blank_streak = 0
        marks = [code for code, col in role_cols.items() if is_marked(ws.cell(row=row, column=col).value)]
        supervised_col = role_cols["staff"] + 1
        entries.append({
            "row": row,
            "id": ident,
            "nama": str(ws.cell(row=row, column=3).value or "").strip(),
            "nip": normalize_digits(ws.cell(row=row, column=4).value),
            "department": str(ws.cell(row=row, column=5).value or "").strip(),
            "marks": marks,
            "supervised": str(ws.cell(row=row, column=supervised_col).value or "").strip(),
        })

    separate: Optional[Dict[str, Any]] = None
    for row in range(1, ws.max_row + 1):
        marker = str(ws.cell(row=row, column=1).value or "")
        if marker.startswith("AKUN SUPER ADMIN TERPISAH"):
            values = [ws.cell(row=row + 2, column=col).value for col in range(1, 5)]
            separate = {
                "name": str(values[0] or "").strip(),
                "nip": normalize_digits(values[1]),
                "pin": normalize_digits(values[2]) or DEFAULT_PIN,
                "note": str(values[3] or "").strip(),
                "row": row + 2,
            }
            break
    return entries, separate


def parse_csv(source: Union[str, Path, bytes]) -> List[Dict[str, Any]]:
    """Baca CSV format clustering (kolom 'ID Akun' + kolom peran 0/1)."""
    if isinstance(source, (bytes, bytearray)):
        text = source.decode("utf-8-sig", errors="replace")
    elif isinstance(source, Path):
        text = source.read_text(encoding="utf-8-sig")
    else:
        text = str(source)

    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise TemplateError("CSV tidak memiliki baris header.")
    fields = {(name or "").strip(): name for name in reader.fieldnames}
    if "ID Akun" not in fields:
        raise TemplateError(
            "Kolom 'ID Akun' tidak ditemukan. Gunakan CSV hasil unduhan menu Clustering akun."
        )
    supervised_key = fields.get("Departemen Diawasi")

    entries: List[Dict[str, Any]] = []
    for index, row in enumerate(reader, start=2):
        ident = str(row.get(fields["ID Akun"]) or "").strip()
        if not ident:
            continue
        marks = [
            code for code, label in ROLES
            if label in fields and is_marked(row.get(fields[label]))
        ]
        entries.append({
            "row": index,
            "id": ident,
            "nama": str(row.get(fields.get("Nama Lengkap", ""), "") or "").strip(),
            "nip": normalize_digits(row.get(fields.get("NIP/NIK", ""), "")),
            "department": str(row.get(fields.get("Departemen", ""), "") or "").strip(),
            "marks": marks,
            "supervised": str(row.get(supervised_key, "") or "").strip() if supervised_key else "",
        })
    if not entries:
        raise TemplateError("CSV tidak memuat baris akun (kolom 'ID Akun' kosong).")
    return entries


def read_upload(filename: str, payload: bytes) -> Tuple[List[Dict[str, Any]], Optional[Dict[str, Any]]]:
    """Pilih pembaca sesuai ekstensi berkas unggahan (.xlsx/.xlsm atau .csv)."""
    if not payload:
        raise TemplateError("Berkas kosong.")
    if len(payload) > MAX_UPLOAD_BYTES:
        raise TemplateError("Ukuran berkas melebihi 8 MB.")
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return parse_csv(payload), None
    if name.endswith((".xlsx", ".xlsm")):
        return parse_workbook(payload)
    # Berkas tanpa ekstensi: coba xlsx (zip) lalu CSV.
    if payload[:2] == b"PK":
        return parse_workbook(payload)
    return parse_csv(payload), None


# --------------------------------------------------------------------------- #
# Menerapkan peran
# --------------------------------------------------------------------------- #
def apply_roles(staff: List[Dict[str, Any]], entries: Iterable[Dict[str, Any]], pin: str,
                strict: bool = False) -> Dict[str, Any]:
    """Terapkan peran + NIP/NIK + departemen diawasi + PIN seragam pada daftar staf.

    Mengubah ``staff`` di tempat (in-place) dan mengembalikan ringkasan hasil.
    Bila ada ``problems``, tidak ada perubahan yang boleh disimpan pemanggil.
    """
    if not (len(str(pin)) == 6 and str(pin).isdigit()):
        raise TemplateError("PIN harus 6 digit angka.")

    entries = list(entries)
    staff_by_id = {str(row.get("id")): row for row in staff}
    seen_nip: Dict[str, str] = {}
    for row in staff:
        nip = normalize_digits(row.get("nip"))
        if nip:
            seen_nip[nip] = str(row.get("id"))

    problems: List[str] = []
    warnings: List[str] = []
    missing_identity: List[Tuple[str, str]] = []
    changed: List[Tuple[str, str, str]] = []
    fields = uniform_pin_fields(pin)

    for entry in entries:
        ident = entry["id"]
        row = staff_by_id.get(ident)
        if row is None:
            problems.append(f"Baris {entry['row']}: ID '{ident}' tidak ada di database.")
            continue
        if row.get("system_super_admin"):
            warnings.append(f"Baris {entry['row']}: '{ident}' adalah akun Super Admin terpisah; dilewati.")
            continue
        if len(entry["marks"]) > 1:
            labels = ", ".join(ROLE_LABEL_BY_CODE[code] for code in entry["marks"])
            problems.append(
                f"Baris {entry['row']} ({ident}): lebih dari satu peran ditandai ({labels}). "
                f"Hapus angka 1 pada kolom peran lain, sisakan satu."
            )
            continue
        if not entry["marks"]:
            if strict:
                problems.append(f"Baris {entry['row']} ({ident}): belum ada peran yang ditandai.")
                continue
            warnings.append(f"Baris {entry['row']} ({ident}): peran kosong, peran lama dipertahankan.")
            role_code = str(row.get("role") or "staff")
        else:
            role_code = entry["marks"][0]

        nip = entry["nip"]
        if nip:
            if len(nip) not in (16, 18):
                problems.append(f"Baris {entry['row']} ({ident}): NIP/NIK '{nip}' bukan 16/18 digit.")
                continue
            owner = seen_nip.get(nip)
            if owner and owner != ident:
                problems.append(f"Baris {entry['row']} ({ident}): NIP/NIK '{nip}' sudah dipakai {owner}.")
                continue
        else:
            nip = normalize_digits(row.get("nip"))
            if not nip:
                missing_identity.append((ident, str(row.get("name") or "")))

        supervised = [
            item.strip() for item in entry["supervised"].replace(",", ";").split(";") if item.strip()
        ]
        if role_code not in SUPERVISOR_ROLES:
            supervised = []
        elif not supervised:
            supervised = [row.get("department")] if row.get("department") else []

        before = str(row.get("role") or "staff")
        row["role"] = role_code
        if nip:
            row["nip"] = nip
            row["id_type"] = id_type_for(nip)
            seen_nip[nip] = ident
        row["supervised_departments"] = supervised
        row.update(fields)
        if before != role_code:
            changed.append((ident, before, role_code))

    return {
        "pin": pin,
        "problems": problems,
        "warnings": warnings,
        "missing_identity": missing_identity,
        "role_changes": changed,
        "processed": len(entries),
    }


def upsert_separate_super_admin(staff: List[Dict[str, Any]], info: Optional[Dict[str, Any]],
                                pin: str = DEFAULT_PIN) -> Dict[str, Any]:
    """Buat/perbarui akun Super Admin terpisah dari blok pada berkas unggahan."""
    if not info or not info.get("nip"):
        return {"created": False, "reason": "NIP/NIK Super Admin terpisah kosong — dilewati."}
    nip = normalize_digits(info["nip"])
    if len(nip) not in (16, 18):
        return {"created": False, "reason": f"NIP/NIK Super Admin '{nip}' bukan 16/18 digit — dilewati."}
    for row in staff:
        if str(row.get("id")) != SEPARATE_SUPER_ADMIN_ID and normalize_digits(row.get("nip")) == nip:
            return {"created": False, "reason": f"NIP/NIK {nip} sudah dipakai akun {row.get('id')}."}

    name = info.get("name") or SEPARATE_SUPER_ADMIN_DEFAULT["name"]
    super_admin_pin = info.get("pin") or pin or DEFAULT_PIN
    existing = separate_super_admin(staff)
    payload = {
        "id": SEPARATE_SUPER_ADMIN_ID,
        "name": name,
        "initials": "".join(part[0] for part in name.split()[:2]).upper() or "SA",
        "department": "Admin",
        "position": "Super Admin",
        "keterangan": "Akun sistem terpisah (di luar roster staf)",
        "active": True,
        "archived": False,
        "nip": nip,
        "id_type": id_type_for(nip),
        "role": "super_admin",
        "system_super_admin": True,
        "supervised_departments": [],
        "identity_is_placeholder": True,
        "credential_source": "clustering-setup",
    }
    payload.update(uniform_pin_fields(super_admin_pin, source="clustering-setup"))
    if existing is None:
        staff.append(payload)
        return {"created": True, "id": SEPARATE_SUPER_ADMIN_ID, "pin": super_admin_pin, "nip": nip}
    existing.update(payload)
    return {"created": False, "updated": True, "id": SEPARATE_SUPER_ADMIN_ID,
            "pin": super_admin_pin, "nip": nip}


def result_payload(result: Dict[str, Any], super_admin: Dict[str, Any]) -> Dict[str, Any]:
    """Ringkasan JSON yang aman dikirim ke klien (tanpa hash PIN)."""
    return {
        "ok": True,
        "pin": result["pin"],
        "processed": result["processed"],
        "role_changes": [
            {"id": ident, "from": before, "from_label": ROLE_LABEL_BY_CODE.get(before, before),
             "to": after, "to_label": ROLE_LABEL_BY_CODE.get(after, after)}
            for ident, before, after in result["role_changes"]
        ],
        "warnings": result["warnings"][:50],
        "warning_count": len(result["warnings"]),
        "missing_identity": [
            {"id": ident, "name": name} for ident, name in result["missing_identity"]
        ],
        "super_admin": super_admin,
    }
