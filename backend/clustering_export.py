"""Ekspor format clustering akun LOKA-Kin (Excel, CSV, dan tampilan cetak).

Modul ini dipakai dua jalur:

1. ``scripts/clustering.py`` untuk membuat berkas isian peran di mesin lokal.
2. Endpoint ``/api/admin/clustering/*`` agar Super Admin dapat mengunduh format
   yang sama langsung dari aplikasi (lihat ``server.py``).

Semua fungsi di sini hanya membaca daftar dokumen staf (dict) dan tidak menyentuh
database, sehingga aman dipanggil dari proses server yang sedang berjalan.
"""

from __future__ import annotations

import csv
import io
from collections import Counter
from html import escape
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

DEFAULT_PIN = "123456"

# Urutan kolom peran pada berkas Excel = urutan hierarki akses.
ROLES: List[Tuple[str, str]] = [
    ("super_admin", "Super Admin"),
    ("head", "Kepala Loka"),
    ("clinical_supervisor", "Clinical Supervisor"),
    ("team_lead", "Ketua Tim"),
    ("staff", "Staf"),
]
ROLE_CODE_BY_LABEL = {label: code for code, label in ROLES}
ROLE_LABEL_BY_CODE = {code: label for code, label in ROLES}
SUPERVISOR_ROLES = {"head", "clinical_supervisor", "team_lead"}

ROLE_ACCESS = {
    "super_admin": "Mengelola seluruh aplikasi: roster, peran, laporan, dan pengaturan website.",
    "head": "Membaca seluruh laporan; membuat/mengubah laporan pada departemen yang diawasi.",
    "clinical_supervisor": "Membaca seluruh laporan; membuat/mengubah laporan pada departemen yang diawasi.",
    "team_lead": "Membaca seluruh laporan; membuat/mengubah laporan pada departemen yang diawasi.",
    "staff": "Hanya membaca dan mengelola laporan miliknya sendiri.",
}

DEPARTMENTS = [
    "Admin", "Bendahara", "Perencanaan", "Informasi dan Humas",
    "Layanan Rehabilitasi Medis", "Layanan Rehabilitasi Sosial", "Umum",
    "Sarana & Prasarana", "Clinical Supervisor",
]

SEPARATE_SUPER_ADMIN_ID = "super-admin"
SEPARATE_SUPER_ADMIN_DEFAULT = {
    "name": "Super Admin Loka-Kin",
    "nip": "0000000000000001",
    "pin": DEFAULT_PIN,
    "note": "Akun sistem terpisah (bukan bagian 78 staf). Ganti NIP/NIK dengan identitas resmi bila sudah tersedia.",
}

COLUMNS = [
    "No", "ID Akun", "Nama Lengkap", "NIP/NIK", "Departemen", "Peran Saat Ini",
    "Super Admin", "Kepala Loka", "Clinical Supervisor", "Ketua Tim", "Staf",
    "Departemen Diawasi", "Catatan",
]


def normalize_digits(value: Any) -> str:
    return "".join(ch for ch in str(value or "") if ch.isdigit())


def role_label(code: str) -> str:
    return ROLE_LABEL_BY_CODE.get(str(code or "staff"), "Staf")


def sort_staff(rows: Iterable[Dict[str, Any]]) -> List[Dict[str, Any]]:
    def key(row: Dict[str, Any]) -> Tuple[int, str]:
        ident = str(row.get("id") or "")
        suffix = ident.rsplit("-", 1)[-1]
        return (int(suffix) if suffix.isdigit() else 10**6, ident)

    return sorted(rows, key=key)


def roster_rows(staff: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Akun roster (78 staf), tanpa akun sistem Super Admin terpisah."""
    rows = [
        row for row in staff
        if row.get("id") and not row.get("system_super_admin") and row.get("id") != SEPARATE_SUPER_ADMIN_ID
    ]
    return sort_staff(rows)


def separate_super_admin(staff: Sequence[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    for row in staff:
        if row.get("system_super_admin") or row.get("id") == SEPARATE_SUPER_ADMIN_ID:
            return row
    return None


def account_rows(staff: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Baris siap tulis untuk sheet 'Isi Peran' (peran terisi dari database)."""
    entries: List[Dict[str, Any]] = []
    for index, row in enumerate(roster_rows(staff), start=1):
        nip = normalize_digits(row.get("nip"))
        role_code = str(row.get("role") or "staff")
        notes: List[str] = []
        if role_code == "super_admin":
            # Super Admin dipindahkan ke akun terpisah, jadi akun roster ini
            # ditandai Staf dan dapat diubah pengguna bila perlu.
            role_code = "staff"
            notes.append("Sebelumnya Super Admin; Super Admin kini akun terpisah — ubah bila perlu.")
        if not nip:
            notes.append("NIP/NIK belum ada — wajib diisi agar akun dapat login.")
        if role_code not in ROLE_LABEL_BY_CODE:
            role_code = "staff"
        entries.append({
            "no": index,
            "id": row.get("id"),
            "name": row.get("name") or "",
            "nip": nip,
            "department": row.get("department") or "",
            "current_label": role_label(row.get("role")),
            "role_code": role_code,
            "supervised": "; ".join(row.get("supervised_departments") or []),
            "note": " ".join(notes),
        })
    return entries


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def build_workbook(staff: Sequence[Dict[str, Any]], pin: str = DEFAULT_PIN):
    """Bangun workbook format clustering (4 sheet) dari daftar dokumen staf."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    entries = account_rows(staff)
    separate = separate_super_admin(staff)

    header_fill = PatternFill("solid", fgColor="1F3864")
    role_fill = PatternFill("solid", fgColor="2E75B6")
    info_fill = PatternFill("solid", fgColor="E8F0FF")
    input_fill = PatternFill("solid", fgColor="FFF7E0")
    title_font = Font(bold=True, size=14, color="1F3864")
    head_font = Font(bold=True, color="FFFFFF", size=10)
    thin = Side(style="thin", color="B4C6E7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()

    # ---------------- Sheet 1: isian peran -------------------------------- #
    ws = wb.active
    ws.title = "Isi Peran"
    ws.sheet_properties.tabColor = "2E75B6"

    last_col = len(COLUMNS)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(row=1, column=1, value="FORMAT CLUSTERING AKUN LOKA-KIN").font = title_font
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    text = ws.cell(
        row=2, column=1,
        value=(
            f"Kolom peran (G–K) sudah terisi otomatis sesuai peran di database — perbaiki bila ada yang keliru, "
            f"lalu jalankan `python scripts/clustering.py apply Clustering-Akun-LokaKin.xlsx`. "
            f"Isi angka 1 pada TEPAT SATU kolom peran per baris. PIN seluruh akun: {pin}."
        ),
    )
    text.alignment = wrap
    text.fill = info_fill

    header_row = 4
    for col, title in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = head_font
        cell.fill = role_fill if 7 <= col <= 11 else header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    data_row = header_row + 1
    for entry in entries:
        marks = [1 if code == entry["role_code"] else None for code, _label in ROLES]
        values = [
            entry["no"], entry["id"], entry["name"], entry["nip"] or None,
            entry["department"], entry["current_label"], *marks,
            entry["supervised"], entry["note"],
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=data_row, column=col, value=value)
            cell.border = border
            cell.alignment = wrap
            if 7 <= col <= 11:
                cell.fill = input_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
        data_row += 1

    first_data, last_data = header_row + 1, data_row - 1
    dv = DataValidation(type="list", formula1='"1"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Nilai tidak valid",
                        error="Isi dengan angka 1 pada satu kolom peran, atau biarkan kosong.")
    dv.errorStyle = "stop"
    ws.add_data_validation(dv)
    dv.add(f"G{first_data}:K{last_data}")

    # Baris dengan lebih dari satu peran ditandai merah agar mudah diperbaiki.
    ws.conditional_formatting.add(
        f"G{first_data}:K{last_data}",
        FormulaRule(formula=[f'COUNTIF($G{first_data}:$K{first_data},"1")>1'],
                    fill=PatternFill("solid", fgColor="FFC7CE"), stopIfTrue=False),
    )

    widths = {"A": 5, "B": 12, "C": 34, "D": 20, "E": 26, "F": 16,
              "G": 13, "H": 14, "I": 19, "J": 12, "K": 9, "L": 30, "M": 46}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.row_dimensions[header_row].height = 34
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{last_data}"

    # ---------------- Blok Super Admin terpisah --------------------------- #
    block_label_row = last_data + 3
    ws.merge_cells(start_row=block_label_row, start_column=1, end_row=block_label_row, end_column=last_col)
    label = ws.cell(row=block_label_row, column=1,
                    value="AKUN SUPER ADMIN TERPISAH (DI LUAR DAFTAR 78 AKUN DI ATAS)")
    label.font = Font(bold=True, size=12, color="FFFFFF")
    label.fill = PatternFill("solid", fgColor="7F6000")
    label.alignment = Alignment(horizontal="left", vertical="center")

    for col, title in enumerate(["Nama Lengkap", "NIP/NIK", "PIN", "Catatan"], start=1):
        cell = ws.cell(row=block_label_row + 1, column=col, value=title)
        cell.font = head_font
        cell.fill = PatternFill("solid", fgColor="BF9000")
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    prefilled = [
        (separate or {}).get("name") or SEPARATE_SUPER_ADMIN_DEFAULT["name"],
        normalize_digits((separate or {}).get("nip")) or SEPARATE_SUPER_ADMIN_DEFAULT["nip"],
        pin,
        SEPARATE_SUPER_ADMIN_DEFAULT["note"],
    ]
    for col, value in enumerate(prefilled, start=1):
        cell = ws.cell(row=block_label_row + 2, column=col, value=value)
        cell.fill = input_fill
        cell.border = border
        cell.alignment = wrap

    ws.merge_cells(start_row=block_label_row + 4, start_column=1,
                   end_row=block_label_row + 4, end_column=last_col)
    hint = ws.cell(
        row=block_label_row + 4, column=1,
        value=("Nama dan NIP/NIK boleh diubah sebelum apply. Kosongkan NIP/NIK untuk melewati pembuatan "
               "akun Super Admin terpisah. Identitas awal ditandai placeholder agar tidak dipakai pada "
               "PDF kredensial pegawai."),
    )
    hint.alignment = wrap
    hint.fill = info_fill

    # ---------------- Sheet 2: snapshot database -------------------------- #
    snap = wb.create_sheet("Snapshot Database")
    snap_headers = ["ID Akun", "Nama Lengkap", "NIP/NIK", "Kode Peran", "Label Peran",
                    "Departemen", "Departemen Diawasi", "Status Aktivasi",
                    "Sumber Kredensial", "Login Terakhir"]
    for col, title in enumerate(snap_headers, start=1):
        cell = snap.cell(row=1, column=col, value=title)
        cell.font = head_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, row in enumerate(roster_rows(staff), start=2):
        values = [
            row.get("id"), row.get("name"), normalize_digits(row.get("nip")),
            row.get("role") or "staff", role_label(row.get("role")),
            row.get("department") or "",
            "; ".join(row.get("supervised_departments") or []),
            "Aktif" if row.get("is_activated") else "Belum aktivasi",
            row.get("credential_source") or "-",
            str(row.get("last_login_at") or "-")[:19],
        ]
        for col, value in enumerate(values, start=1):
            cell = snap.cell(row=index, column=col, value=value)
            cell.border = border
            cell.alignment = wrap
    for col in range(1, len(snap_headers) + 1):
        snap.column_dimensions[get_column_letter(col)].width = 24 if col > 2 else 14
    snap.freeze_panes = "A2"
    snap.auto_filter.ref = f"A1:J{max(1, len(entries) + 1)}"

    # ---------------- Sheet 3: matriks cluster ---------------------------- #
    matrix = wb.create_sheet("Matriks Cluster")
    matrix.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + len(ROLES))
    matrix.cell(row=1, column=1, value="MATRIKS CLUSTER: JUMLAH AKUN PER DEPARTEMEN x PERAN").font = title_font
    for col, title in enumerate(["Departemen"] + [label for _code, label in ROLES] + ["Total"], start=1):
        cell = matrix.cell(row=3, column=col, value=title)
        cell.font = head_font
        cell.fill = role_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    summary = summarize(staff)
    excel_row = 4
    for department in summary["departments"]:
        matrix.cell(row=excel_row, column=1, value=department).border = border
        for offset, (code, _label) in enumerate(ROLES, start=2):
            cell = matrix.cell(row=excel_row, column=offset,
                               value=summary["matrix"].get(department, {}).get(code, 0))
            cell.border = border
            cell.alignment = Alignment(horizontal="center")
        total = matrix.cell(row=excel_row, column=2 + len(ROLES),
                            value=summary["row_totals"].get(department, 0))
        total.border = border
        total.alignment = Alignment(horizontal="center")
        excel_row += 1

    matrix.cell(row=excel_row, column=1, value="TOTAL").font = Font(bold=True)
    for offset, (code, _label) in enumerate(ROLES, start=2):
        cell = matrix.cell(row=excel_row, column=offset, value=summary["role_counts"].get(code, 0))
        cell.font = Font(bold=True)
        cell.border = border
        cell.alignment = Alignment(horizontal="center")
    total_all = matrix.cell(row=excel_row, column=2 + len(ROLES), value=summary["roster_total"])
    total_all.font = Font(bold=True)
    matrix.column_dimensions["A"].width = 32
    for col in range(2, 3 + len(ROLES)):
        matrix.column_dimensions[get_column_letter(col)].width = 18

    # ---------------- Sheet 4: panduan ------------------------------------ #
    guide = wb.create_sheet("Panduan")
    guide.column_dimensions["A"].width = 26
    guide.column_dimensions["B"].width = 96
    guide.cell(row=1, column=1, value="PANDUAN PENGGUNAAN").font = title_font
    steps = [
        ("1. Isi peran", "Buka sheet 'Isi Peran'. Isi angka 1 pada tepat satu kolom peran (Super Admin / "
                         "Kepala Loka / Clinical Supervisor / Ketua Tim / Staf). Kolom sudah terisi otomatis; "
                         "saat memindahkan peran, hapus dulu angka 1 lama (pilih sel, tekan Delete) agar tidak "
                         "ada dua peran pada satu baris — baris dengan dua peran ditandai merah."),
        ("2. Lengkapi NIP/NIK", "Login memakai NIP/NIK 16 atau 18 digit. Baris dengan NIP/NIK kosong berstatus "
                                "'Belum aktivasi' dan tidak dapat dipakai clustering sampai kolom D diisi."),
        ("3. Departemen diawasi", "Wajib untuk Kepala Loka, Clinical Supervisor, dan Ketua Tim. Pisahkan beberapa "
                                  "departemen dengan titik koma. Bila kosong, sistem memakai departemen akun itu sendiri."),
        ("4. Super Admin", "Isi blok di bawah tabel untuk akun Super Admin terpisah (bukan bagian 78 akun). "
                           "Kosongkan NIP/NIK bila belum ingin membuatnya."),
        ("5. Terapkan", "Hentikan server, jalankan `python scripts/clustering.py apply Clustering-Akun-LokaKin.xlsx`, "
                        "lalu nyalakan server kembali. Perintah ini juga menyeragamkan PIN."),
        ("6. Periksa", "Jalankan `python scripts/clustering.py report` untuk melihat matriks cluster dan "
                       "daftar akun yang belum lengkap."),
        ("Keamanan", f"PIN seragam {pin} hanya untuk keperluan clustering/pelatihan. Setelah selesai, rotasi PIN "
                     "lewat menu Pengaturan website → Kredensial pegawai atau minta setiap pengguna mengganti PIN."),
    ]
    current = 3
    for title, body in steps:
        cell = guide.cell(row=current, column=1, value=title)
        cell.font = Font(bold=True, color="1F3864")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        body_cell = guide.cell(row=current, column=2, value=body)
        body_cell.alignment = Alignment(vertical="top", wrap_text=True)
        current += 1

    current += 1
    guide.cell(row=current, column=1, value="KODE PERAN & HAK AKSES").font = Font(bold=True, size=12, color="1F3864")
    current += 1
    for col, title in enumerate(["Label pada Excel", "Kode database / hak akses"], start=1):
        cell = guide.cell(row=current, column=col, value=title)
        cell.font = head_font
        cell.fill = header_fill
        cell.border = border
    current += 1
    for code, label_text in ROLES:
        guide.cell(row=current, column=1, value=label_text).border = border
        cell = guide.cell(row=current, column=2, value=f"{code} — {ROLE_ACCESS[code]}")
        cell.border = border
        cell.alignment = wrap
        current += 1

    current += 1
    guide.cell(row=current, column=1, value="DEPARTEMEN RESMI").font = Font(bold=True, size=12, color="1F3864")
    current += 1
    guide.cell(row=current, column=2, value="; ".join(DEPARTMENTS)).alignment = wrap

    return wb


# --------------------------------------------------------------------------- #
# CSV & HTML
# --------------------------------------------------------------------------- #
def build_csv(staff: Sequence[Dict[str, Any]], pin: str = DEFAULT_PIN) -> bytes:
    """CSV datar 78 akun: satu kolom per peran, siap dibaca pandas/SPSS/Excel."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(COLUMNS)
    for entry in account_rows(staff):
        writer.writerow([
            entry["no"], entry["id"], entry["name"], entry["nip"], entry["department"],
            entry["current_label"],
            *[1 if code == entry["role_code"] else 0 for code, _label in ROLES],
            entry["supervised"], entry["note"],
        ])
    # BOM agar Excel membaca UTF-8 dengan benar.
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def build_print_html(staff: Sequence[Dict[str, Any]], pin: str = DEFAULT_PIN,
                     title: str = "Format Clustering Akun LOKA-Kin") -> str:
    """Halaman HTML siap cetak (A4 lanskap) sebagai alternatif PDF/Excel."""
    entries = account_rows(staff)
    separate = separate_super_admin(staff)
    summary = summarize(staff)
    role_headers = "".join(f"<th>{escape(label)}</th>" for _code, label in ROLES)
    body_rows = []
    for entry in entries:
        marks = "".join(
            f"<td class='mark'>{'1' if code == entry['role_code'] else ''}</td>" for code, _label in ROLES
        )
        body_rows.append(
            "<tr>"
            f"<td class='num'>{entry['no']}</td>"
            f"<td>{escape(str(entry['id'] or ''))}</td>"
            f"<td>{escape(entry['name'])}</td>"
            f"<td class='mono'>{escape(entry['nip'] or '—')}</td>"
            f"<td>{escape(entry['department'])}</td>"
            f"<td>{escape(entry['current_label'])}</td>"
            f"{marks}"
            f"<td>{escape(entry['supervised'] or '—')}</td>"
            "</tr>"
        )

    matrix_rows = []
    for department in summary["departments"]:
        cells = "".join(
            f"<td class='num'>{summary['matrix'].get(department, {}).get(code, 0) or '—'}</td>"
            for code, _label in ROLES
        )
        matrix_rows.append(
            f"<tr><td>{escape(department)}</td>{cells}"
            f"<td class='num'><b>{summary['row_totals'].get(department, 0)}</b></td></tr>"
        )
    total_cells = "".join(
        f"<td class='num'><b>{summary['role_counts'].get(code, 0)}</b></td>" for code, _label in ROLES
    )

    super_admin_block = ""
    if separate:
        super_admin_block = (
            "<section class='block'>"
            "<h2>Akun Super Admin Terpisah (di luar 78 akun)</h2>"
            "<table class='grid'><tr><th>Nama</th><th>NIP/NIK</th><th>PIN</th></tr>"
            f"<tr><td>{escape(str(separate.get('name') or ''))}</td>"
            f"<td class='mono'>{escape(normalize_digits(separate.get('nip')) or '—')}</td>"
            f"<td class='mono'>{escape(pin)}</td></tr></table></section>"
        )

    return f"""<!doctype html>
<html lang="id"><head><meta charset="utf-8">
<title>{escape(title)}</title>
<style>
  @page {{ size: A4 landscape; margin: 12mm; }}
  * {{ box-sizing: border-box; }}
  body {{ font-family: "Segoe UI", Arial, sans-serif; color: #1f2937; font-size: 9pt; }}
  h1 {{ font-size: 15pt; color: #1f3864; margin: 0 0 4px; }}
  h2 {{ font-size: 11pt; color: #1f3864; margin: 14px 0 6px; }}
  p.meta {{ margin: 0 0 10px; color: #475569; }}
  table {{ border-collapse: collapse; width: 100%; }}
  th, td {{ border: 1px solid #b4c6e7; padding: 3px 4px; vertical-align: top; }}
  thead th {{ background: #1f3864; color: #fff; font-size: 8pt; text-align: center; }}
  thead th.role {{ background: #2e75b6; }}
  tbody tr:nth-child(even) {{ background: #f6f9ff; }}
  td.num, td.mark {{ text-align: center; }}
  td.mono {{ font-family: "Consolas", monospace; letter-spacing: .3px; }}
  .mark {{ font-weight: 700; color: #1d4ed8; }}
  .note {{ margin-top: 10px; font-size: 8pt; color: #475569; }}
  .toolbar {{ margin-bottom: 10px; }}
  button {{ padding: 7px 14px; border: 0; border-radius: 6px; background: #2458ec; color: #fff; font-weight: 600; cursor: pointer; }}
  @media print {{ .toolbar {{ display: none; }} }}
</style></head>
<body>
<div class="toolbar"><button type="button" onclick="window.print()">Cetak / Simpan sebagai PDF</button></div>
<h1>{escape(title)}</h1>
<p class="meta">Dicetak {escape(_now_text())} · PIN seragam <b>{escape(pin)}</b> · {summary['roster_total']} akun roster
 · {summary['missing_identity']} akun belum punya NIP/NIK</p>
<table class="grid">
  <thead><tr><th>No</th><th>ID Akun</th><th>Nama Lengkap</th><th>NIP/NIK</th><th>Departemen</th>
  <th>Peran Saat Ini</th>{role_headers}<th>Departemen Diawasi</th></tr></thead>
  <tbody>{''.join(body_rows)}</tbody>
</table>
{super_admin_block}
<h2>Matriks Cluster: Departemen × Peran</h2>
<table class="grid">
  <thead><tr><th>Departemen</th>{role_headers}<th>Total</th></tr></thead>
  <tbody>{''.join(matrix_rows)}
  <tr><td><b>TOTAL</b></td>{total_cells}<td class="num"><b>{summary['roster_total']}</b></td></tr></tbody>
</table>
<p class="note">Isi kolom peran pada berkas Excel (angka 1 pada satu kolom) lalu jalankan
<code>python scripts/clustering.py apply Clustering-Akun-LokaKin.xlsx</code> untuk menerapkan peran dan PIN seragam.
Kolom peran di atas tercetak sebagai informasi peran yang tersimpan saat ini.</p>
</body></html>"""


def build_pdf(staff: Sequence[Dict[str, Any]], pin: str = DEFAULT_PIN,
              title: str = "Format Clustering Akun LOKA-Kin") -> bytes:
    """PDF A4 lanskap berisi tabel akun + matriks cluster (siap dicetak)."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    entries = account_rows(staff)
    summary = summarize(staff)
    separate = separate_super_admin(staff)

    buffer = io.BytesIO()
    page_width, page_height = landscape(A4)
    margin = 10 * mm
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=12 * mm,
        title=title, author="LOKA-Kin",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Judul", parent=styles["Title"], fontName="Helvetica-Bold",
                                 fontSize=15, leading=18, alignment=TA_LEFT,
                                 textColor=colors.HexColor("#1F3864"), spaceAfter=2)
    meta_style = ParagraphStyle("Meta", parent=styles["Normal"], fontName="Helvetica",
                                fontSize=7.5, leading=10, textColor=colors.HexColor("#475569"))
    cell_style = ParagraphStyle("Sel", parent=styles["Normal"], fontName="Helvetica",
                                fontSize=6.4, leading=7.6)
    head_style = ParagraphStyle("Kepala", parent=cell_style, fontName="Helvetica-Bold",
                                textColor=colors.white, alignment=1)

    story: List[Any] = [
        Paragraph(escape(title), title_style),
        Paragraph(
            f"Dicetak {escape(_now_text())} &middot; PIN seragam <b>{escape(pin)}</b> &middot; "
            f"{summary['roster_total']} akun roster &middot; {summary['missing_identity']} akun belum punya NIP/NIK",
            meta_style,
        ),
        Spacer(1, 5 * mm),
    ]

    header = ["No", "ID Akun", "Nama Lengkap", "NIP/NIK", "Departemen", "Peran Saat Ini",
              *[label for _code, label in ROLES], "Departemen Diawasi"]
    data: List[List[Any]] = [[Paragraph(escape(h), head_style) for h in header]]
    for entry in entries:
        data.append([
            Paragraph(str(entry["no"]), cell_style),
            Paragraph(escape(str(entry["id"] or "")), cell_style),
            Paragraph(escape(entry["name"]), cell_style),
            Paragraph(escape(entry["nip"] or "—"), cell_style),
            Paragraph(escape(entry["department"]), cell_style),
            Paragraph(escape(entry["current_label"]), cell_style),
            *[Paragraph("1" if code == entry["role_code"] else "", cell_style) for code, _label in ROLES],
            Paragraph(escape(entry["supervised"] or "—"), cell_style),
        ])

    available = page_width - 2 * margin
    ratios = [0.035, 0.075, 0.20, 0.115, 0.155, 0.09, 0.06, 0.065, 0.085, 0.055, 0.04, 0.125]
    widths = [available * ratio for ratio in ratios]
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
        ("BACKGROUND", (6, 0), (10, 0), colors.HexColor("#2E75B6")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B4C6E7")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (0, 1), (1, -1), "CENTER"),
        ("ALIGN", (6, 1), (10, -1), "CENTER"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F9FF")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 2.5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2.5),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    story.append(table)

    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph("Matriks Cluster: Departemen &times; Peran", title_style))
    matrix_head = ["Departemen", *[label for _code, label in ROLES], "Total"]
    matrix_data: List[List[Any]] = [[Paragraph(escape(h), head_style) for h in matrix_head]]
    for department in summary["departments"]:
        cells = [Paragraph(escape(department), cell_style)]
        cells += [Paragraph(str(summary["matrix"][department].get(code, 0)), cell_style) for code, _l in ROLES]
        cells.append(Paragraph(str(summary["row_totals"].get(department, 0)), cell_style))
        matrix_data.append(cells)
    totals_row = [Paragraph("<b>TOTAL</b>", cell_style)]
    totals_row += [Paragraph(f"<b>{summary['role_counts'].get(code, 0)}</b>", cell_style) for code, _l in ROLES]
    totals_row.append(Paragraph(f"<b>{summary['roster_total']}</b>", cell_style))
    matrix_data.append(totals_row)

    matrix = Table(matrix_data, colWidths=[available * 0.26] + [available * 0.11] * len(ROLES) + [available * 0.08],
                   repeatRows=1)
    matrix.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E75B6")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B4C6E7")),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F6F9FF")]),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E8F0FF")),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    story.append(matrix)

    if separate:
        story.append(Spacer(1, 5 * mm))
        story.append(Paragraph("Akun Super Admin Terpisah (di luar 78 akun)", title_style))
        sa = Table([
            [Paragraph(escape(str(separate.get("name") or "")), cell_style),
             Paragraph(escape(normalize_digits(separate.get("nip")) or "—"), cell_style),
             Paragraph(escape(pin), cell_style)],
        ], colWidths=[available * 0.4, available * 0.3, available * 0.15])
        sa.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#B4C6E7")),
                                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B4C6E7"))]))
        story.append(sa)

    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(
        "Isi angka 1 pada satu kolom peran di berkas Excel (menu Super Admin &rarr; Clustering akun) lalu jalankan "
        "<i>clustering.py apply</i> untuk menerapkan peran dan PIN seragam. PIN 123456 hanya untuk "
        "keperluan clustering/pelatihan.",
        meta_style,
    ))

    doc.build(story)
    return buffer.getvalue()


def _now_text() -> str:
    from datetime import datetime, timezone, timedelta

    jakarta = timezone(timedelta(hours=7))
    return datetime.now(jakarta).strftime("%d-%m-%Y %H:%M WIB")


# --------------------------------------------------------------------------- #
# Ringkasan
# --------------------------------------------------------------------------- #
def summarize(staff: Sequence[Dict[str, Any]]) -> Dict[str, Any]:
    rows = roster_rows(staff)
    role_counts: Counter = Counter(str(row.get("role") or "staff") for row in rows)
    departments = sorted({str(row.get("department") or "-") for row in rows})
    matrix: Dict[str, Dict[str, int]] = {}
    row_totals: Dict[str, int] = {}
    for department in departments:
        counts = Counter(
            str(row.get("role") or "staff")
            for row in rows if str(row.get("department") or "-") == department
        )
        matrix[department] = {code: counts.get(code, 0) for code, _label in ROLES}
        row_totals[department] = sum(counts.values())
    missing = [row for row in rows if len(normalize_digits(row.get("nip"))) not in (16, 18)]
    return {
        "roster_total": len(rows),
        "role_counts": {code: role_counts.get(code, 0) for code, _label in ROLES},
        "role_labels": {code: label for code, label in ROLES},
        "departments": departments,
        "matrix": matrix,
        "row_totals": row_totals,
        "missing_identity": len(missing),
        "missing_staff": [
            {"id": row.get("id"), "name": row.get("name")} for row in missing[:12]
        ],
        "separate_super_admin": (
            {
                "id": separate_super_admin(staff).get("id"),
                "name": separate_super_admin(staff).get("name"),
                "nip": normalize_digits(separate_super_admin(staff).get("nip")),
                "active": bool(separate_super_admin(staff).get("is_activated")),
            }
            if separate_super_admin(staff) else None
        ),
    }
