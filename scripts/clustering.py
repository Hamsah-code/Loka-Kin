#!/usr/bin/env python3
"""Clustering akun LOKA-Kin: penyeragaman PIN, format peran, dan laporan cluster.

Perintah
--------
  template   Buat berkas Excel untuk mengisi peran setiap akun (satu kolom per
             peran: Super Admin, Kepala Loka, Clinical Supervisor, Ketua Tim,
             Staf) sekaligus blok akun Super Admin terpisah.
  pin        Seragamkan PIN semua akun roster (default 123456) dan aktifkan
             sehingga setiap akun dapat dipakai untuk clustering.
  apply      Terapkan peran dari berkas Excel ke database (PIN tetap diseragamkan).
  report     Tampilkan matriks cluster (peran x departemen) dari database.

Contoh
------
  .venv/bin/python scripts/clustering.py template
  .venv/bin/python scripts/clustering.py pin --pin 123456
  .venv/bin/python scripts/clustering.py apply Clustering-Akun-LokaKin.xlsx
  .venv/bin/python scripts/clustering.py report

Skrip ini menulis langsung ke snapshot JSON mode ``local-json``
(``backend/data/loka_kin_db.json``) yang dipakai bila ``MONGO_URL`` kosong.
Hentikan server pratinjau sebelum menjalankan ``pin``/``apply`` agar isi memori
server tidak menimpa perubahan: hentikan proses, jalankan perintah, nyalakan lagi.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

REPO = Path(__file__).resolve().parents[1]
BACKEND = REPO / "backend"
DATA_FILE = BACKEND / "data" / "loka_kin_db.json"
DEFAULT_TEMPLATE = REPO / "Clustering-Akun-LokaKin.xlsx"
DEFAULT_PIN = "123456"

sys.path.insert(0, str(BACKEND))

# Urutan kolom peran pada berkas Excel = urutan hierarki akses.
ROLES: List[Tuple[str, str]] = [
    ("super_admin", "Super Admin"),
    ("head", "Kepala Loka"),
    ("clinical_supervisor", "Clinical Supervisor"),
    ("team_lead", "Ketua Tim"),
    ("staff", "Staf"),
]
ROLE_CODE_BY_LABEL = {label: code for code, label in ROLES}
ROLE_LABEL_BY_CODE = {code: label for code, label in ROLES}
SUPERVISOR_ROLES = {"head", "clinical_supervisor", "team_lead"}

# Keterangan singkat hak akses, dipakai pada sheet Panduan.
ROLE_ACCESS = {
    "super_admin": "Mengelola seluruh aplikasi: roster, peran, laporan, dan pengaturan website.",
    "head": "Membaca seluruh laporan; membuat/mengubah laporan pada departemen yang diawasi.",
    "clinical_supervisor": "Membaca seluruh laporan; membuat/mengubah laporan pada departemen yang diawasi.",
    "team_lead": "Membaca seluruh laporan; membuat/mengubah laporan pada departemen yang diawasi.",
    "staff": "Hanya membaca dan mengelola laporan miliknya sendiri.",
}

# Akun Super Admin terpisah (di luar 78 akun roster). Identitas awal bersifat
# sementara dan ditandai placeholder agar tidak dipakai untuk PDF kredensial.
SEPARATE_SUPER_ADMIN_ID = "super-admin"
SEPARATE_SUPER_ADMIN_DEFAULT = {
    "name": "Super Admin Loka-Kin",
    "nip": "0000000000000001",
    "pin": DEFAULT_PIN,
    "note": "Akun sistem terpisah (bukan bagian 78 staf). Ganti NIP/NIK dengan identitas resmi bila sudah tersedia.",
}


# --------------------------------------------------------------------------- #
# Penyimpanan
# --------------------------------------------------------------------------- #
class Snapshot:
    """Snapshot JSON yang sama dengan yang dipakai ``backend/db.py``."""

    def __init__(self, path: Path = DATA_FILE):
        self.path = Path(path)
        if not self.path.is_file():
            raise SystemExit(f"Snapshot database tidak ditemukan: {self.path}")
        raw = json.loads(self.path.read_text(encoding="utf-8"))
        self.data: Dict[str, Any] = raw if isinstance(raw, dict) else {}
        self.data.setdefault("staff", [])

    @property
    def staff(self) -> List[Dict[str, Any]]:
        return self.data["staff"]

    def backup(self) -> Path:
        stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
        target = self.path.with_name(f"{self.path.stem}.backup-{stamp}.json")
        target.write_text(json.dumps(self.data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        return target

    def save(self) -> None:
        content = json.dumps(self.data, ensure_ascii=False, indent=2, default=str)
        fd, tmp_path = tempfile.mkstemp(dir=str(self.path.parent), suffix=".tmp")
        try:
            with os.fdopen(fd, "w", encoding="utf-8") as handle:
                handle.write(content)
            os.replace(tmp_path, self.path)
        except BaseException:
            if os.path.exists(tmp_path):
                os.unlink(tmp_path)
            raise


def roster_staff(store: Snapshot) -> List[Dict[str, Any]]:
    """Akun roster resmi (78 staf), tanpa akun sistem Super Admin terpisah."""
    rows = [
        row for row in store.staff
        if row.get("id") and not row.get("system_super_admin") and row.get("id") != SEPARATE_SUPER_ADMIN_ID
    ]

    def sort_key(row: Dict[str, Any]) -> Tuple[int, str]:
        ident = str(row.get("id", ""))
        suffix = ident.rsplit("-", 1)[-1]
        return (int(suffix) if suffix.isdigit() else 10**6, ident)

    return sorted(rows, key=sort_key)


def account_label(row: Dict[str, Any]) -> str:
    return ROLE_LABEL_BY_CODE.get(str(row.get("role") or "staff"), str(row.get("role") or "Staf"))


def normalize_digits(value: Any) -> str:
    return "".join(ch for ch in str(value or "") if ch.isdigit())


def id_type_for(nip: str) -> str:
    return {16: "NIK", 18: "NIP"}.get(len(nip), "")


def uniform_pin_fields(pin: str, source: str = "uniform-clustering") -> Dict[str, Any]:
    from auth_utils import hash_credential

    return {
        "pin_hash": hash_credential(pin),
        "is_activated": True,
        "pin_change_required": False,
        "activation_code_hash": "",
        "activation_code_expires_at": "",
        "activation_code_attempts": 0,
        "activation_code_locked_until": "",
        "login_failed_attempts": 0,
        "login_locked_until": "",
        "pin_source": source,
        "pin_updated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
    }


# --------------------------------------------------------------------------- #
# template
# --------------------------------------------------------------------------- #
def build_template(store: Snapshot, out_path: Path, pin: str) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    staff = roster_staff(store)
    separate = next(
        (row for row in store.staff
         if row.get("system_super_admin") or row.get("id") == SEPARATE_SUPER_ADMIN_ID),
        None,
    )

    header_fill = PatternFill("solid", fgColor="1F3864")
    role_fill = PatternFill("solid", fgColor="2E75B6")
    info_fill = PatternFill("solid", fgColor="E8F0FF")
    input_fill = PatternFill("solid", fgColor="FFF7E0")
    title_font = Font(bold=True, size=14, color="1F3864")
    head_font = Font(bold=True, color="FFFFFF", size=10)
    thin = Side(style="thin", color="B4C6E7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()

    # ---------------- Sheet 1: isian peran -------------------------------- #
    ws = wb.active
    ws.title = "Isi Peran"
    ws.sheet_properties.tabColor = "2E75B6"

    headers = [
        "No", "ID Akun", "Nama Lengkap", "NIP/NIK", "Departemen", "Peran Saat Ini",
        "Super Admin", "Kepala Loka", "Clinical Supervisor", "Ketua Tim", "Staf",
        "Departemen Diawasi", "Catatan",
    ]
    last_col = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(row=1, column=1, value="FORMAT CLUSTERING AKUN LOKA-KIN").font = title_font
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    text = ws.cell(
        row=2, column=1,
        value=(
            f"Kolom peran (G–K) sudah terisi otomatis sesuai peran di database — perbaiki bila ada yang keliru, "
            f"lalu jalankan `python scripts/clustering.py apply Clustering-Akun-LokaKin.xlsx`. "
            f"Isi angka 1 pada TEPAT SATU kolom peran per baris. PIN seluruh akun: {pin}."
        ),
    )
    text.alignment = wrap
    text.fill = info_fill

    header_row = 4
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = head_font
        cell.fill = role_fill if 7 <= col <= 11 else header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    data_row = header_row + 1
    for index, row in enumerate(staff, start=1):
        nip = normalize_digits(row.get("nip"))
        department = row.get("department") or ""
        role_code = str(row.get("role") or "staff")
        notes = []
        if role_code == "super_admin":
            # Super Admin dipindahkan ke akun terpisah, jadi akun roster ini
            # ditandai Staf dan dapat diubah pengguna bila perlu.
            role_code = "staff"
            notes.append("Sebelumnya Super Admin; Super Admin kini akun terpisah — ubah bila perlu.")
        if not nip:
            notes.append("NIP/NIK belum ada — wajib diisi agar akun dapat login.")
        marks = [1 if code == role_code else None for code, _label in ROLES]
        values = [
            index, row.get("id"), row.get("name"), nip, department, account_label(row),
            *marks,
            "; ".join(row.get("supervised_departments") or []), " ".join(notes),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=data_row, column=col, value=value)
            cell.border = border
            cell.alignment = wrap
            if 7 <= col <= 11:
                cell.fill = input_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col in (1, 2, 4):
                cell.alignment = Alignment(horizontal="left", vertical="center")
        data_row += 1

    first_data, last_data = header_row + 1, data_row - 1
    dv = DataValidation(type="list", formula1='"1"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Nilai tidak valid",
                        error="Isi dengan angka 1 pada satu kolom peran, atau biarkan kosong.")
    dv.errorStyle = "stop"
    ws.add_data_validation(dv)
    dv.add(f"G{first_data}:K{last_data}")

    # Baris dengan lebih dari satu peran ditandai merah agar mudah diperbaiki.
    from openpyxl.formatting.rule import FormulaRule

    double_mark = PatternFill("solid", fgColor="FFC7CE")
    ws.conditional_formatting.add(
        f"G{first_data}:K{last_data}",
        FormulaRule(formula=[f'COUNTIF($G{first_data}:$K{first_data},"1")>1'],
                    fill=double_mark, stopIfTrue=False),
    )

    widths = {"A": 5, "B": 12, "C": 34, "D": 20, "E": 26, "F": 16,
              "G": 13, "H": 14, "I": 19, "J": 12, "K": 9, "L": 30, "M": 46}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.row_dimensions[header_row].height = 34
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{last_data}"

    # ---------------- Blok Super Admin terpisah --------------------------- #
    block_label_row = last_data + 3
    ws.merge_cells(start_row=block_label_row, start_column=1, end_row=block_label_row, end_column=last_col)
    label = ws.cell(row=block_label_row, column=1,
                    value="AKUN SUPER ADMIN TERPISAH (DI LUAR DAFTAR 78 AKUN DI ATAS)")
    label.font = Font(bold=True, size=12, color="FFFFFF")
    label.fill = PatternFill("solid", fgColor="7F6000")
    label.alignment = Alignment(horizontal="left", vertical="center")

    sub_headers = ["Nama Lengkap", "NIP/NIK", "PIN", "Catatan"]
    for col, title in enumerate(sub_headers, start=1):
        cell = ws.cell(row=block_label_row + 1, column=col, value=title)
        cell.font = head_font
        cell.fill = PatternFill("solid", fgColor="BF9000")
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    prefilled_name = (separate or {}).get("name") or SEPARATE_SUPER_ADMIN_DEFAULT["name"]
    prefilled_nip = normalize_digits((separate or {}).get("nip")) or SEPARATE_SUPER_ADMIN_DEFAULT["nip"]
    prefilled_note = SEPARATE_SUPER_ADMIN_DEFAULT["note"]
    for col, value in enumerate([prefilled_name, prefilled_nip, pin, prefilled_note], start=1):
        cell = ws.cell(row=block_label_row + 2, column=col, value=value)
        cell.fill = input_fill
        cell.border = border
        cell.alignment = wrap

    ws.merge_cells(start_row=block_label_row + 4, start_column=1,
                   end_row=block_label_row + 4, end_column=last_col)
    hint = ws.cell(
        row=block_label_row + 4, column=1,
        value=("Nama dan NIP/NIK boleh diubah sebelum apply. Kosongkan NIP/NIK untuk melewati pembuatan "
               "akun Super Admin terpisah. Identitas awal ditandai placeholder agar tidak dipakai pada "
               "PDF kredensial pegawai."),
    )
    hint.alignment = wrap
    hint.fill = info_fill

    # ---------------- Sheet 2: snapshot database -------------------------- #
    snap = wb.create_sheet("Snapshot Database")
    snap_headers = ["ID Akun", "Nama Lengkap", "NIP/NIK", "Kode Peran", "Label Peran",
                    "Departemen", "Departemen Diawasi", "Status Aktivasi",
                    "Sumber Kredensial", "Login Terakhir"]
    for col, title in enumerate(snap_headers, start=1):
        cell = snap.cell(row=1, column=col, value=title)
        cell.font = head_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, row in enumerate(staff, start=2):
        values = [
            row.get("id"), row.get("name"), normalize_digits(row.get("nip")),
            row.get("role") or "staff", account_label(row),
            row.get("department") or "",
            "; ".join(row.get("supervised_departments") or []),
            "Aktif" if row.get("is_activated") else "Belum aktivasi",
            row.get("credential_source") or "-",
            (row.get("last_login_at") or "-")[:19],
        ]
        for col, value in enumerate(values, start=1):
            cell = snap.cell(row=index, column=col, value=value)
            cell.border = border
            cell.alignment = wrap
    for col in range(1, len(snap_headers) + 1):
        snap.column_dimensions[get_column_letter(col)].width = 24 if col > 2 else 14
    snap.freeze_panes = "A2"
    snap.auto_filter.ref = f"A1:J{max(1, len(staff) + 1)}"

    # ---------------- Sheet 3: matriks cluster ---------------------------- #
    matrix = wb.create_sheet("Matriks Cluster")
    matrix.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + len(ROLES))
    matrix.cell(row=1, column=1, value="MATRIKS CLUSTER: JUMLAH AKUN PER DEPARTEMEN x PERAN").font = title_font
    matrix_head = ["Departemen"] + [label for _code, label in ROLES] + ["Total"]
    for col, title in enumerate(matrix_head, start=1):
        cell = matrix.cell(row=3, column=col, value=title)
        cell.font = head_font
        cell.fill = role_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    counts: Counter = Counter()
    for row in staff:
        counts[(row.get("department") or "-", str(row.get("role") or "staff"))] += 1
    departments = sorted({key[0] for key in counts})
    excel_row = 4
    for department in departments:
        matrix.cell(row=excel_row, column=1, value=department).border = border
        for offset, (code, _label) in enumerate(ROLES, start=2):
            cell = matrix.cell(row=excel_row, column=offset, value=counts.get((department, code), 0))
            cell.border = border
            cell.alignment = Alignment(horizontal="center")
        total = matrix.cell(row=excel_row, column=2 + len(ROLES),
                            value=sum(counts.get((department, code), 0) for code, _ in ROLES))
        total.border = border
        total.alignment = Alignment(horizontal="center")
        excel_row += 1
    matrix.cell(row=excel_row, column=1, value="TOTAL").font = Font(bold=True)
    for offset, (code, _label) in enumerate(ROLES, start=2):
        cell = matrix.cell(row=excel_row, column=offset,
                           value=sum(counts.get((department, code), 0) for department in departments))
        cell.font = Font(bold=True)
        cell.border = border
        cell.alignment = Alignment(horizontal="center")
    total_all = sum(counts.values())
    matrix.cell(row=excel_row, column=2 + len(ROLES), value=total_all).font = Font(bold=True)
    matrix.column_dimensions["A"].width = 32
    for col in range(2, 3 + len(ROLES)):
        matrix.column_dimensions[get_column_letter(col)].width = 18

    # ---------------- Sheet 4: panduan ------------------------------------ #
    guide = wb.create_sheet("Panduan")
    guide.column_dimensions["A"].width = 26
    guide.column_dimensions["B"].width = 96
    guide.cell(row=1, column=1, value="PANDUAN PENGGUNAAN").font = title_font
    steps = [
        ("1. Isi peran", "Buka sheet 'Isi Peran'. Isi angka 1 pada tepat satu kolom peran (Super Admin / "
                         "Kepala Loka / Clinical Supervisor / Ketua Tim / Staf). Kolom sudah terisi otomatis; "
                         "saat memindahkan peran, hapus dulu angka 1 lama (pilih sel, tekan Delete) agar tidak "
                         "ada dua peran pada satu baris — baris dengan dua peran ditandai merah."),
        ("2. Lengkapi NIP/NIK", "Login memakai NIP/NIK 16 atau 18 digit. Baris dengan NIP/NIK kosong berstatus "
                                "'Belum aktivasi' dan tidak dapat dipakai clustering sampai kolom D diisi."),
        ("3. Departemen diawasi", "Wajib untuk Kepala Loka, Clinical Supervisor, dan Ketua Tim. Pisahkan beberapa "
                                  "departemen dengan titik koma. Bila kosong, sistem memakai departemen akun itu sendiri."),
        ("4. Super Admin", "Isi blok di bawah tabel untuk akun Super Admin terpisah (bukan bagian 78 akun). "
                           "Kosongkan NIP/NIK bila belum ingin membuatnya."),
        ("5. Terapkan", "Hentikan server, jalankan `python scripts/clustering.py apply Clustering-Akun-LokaKin.xlsx`, "
                        "lalu nyalakan server kembali. Perintah ini juga menyeragamkan PIN."),
        ("6. Periksa", "Jalankan `python scripts/clustering.py report` untuk melihat matriks cluster dan "
                       "daftar akun yang belum lengkap."),
        ("Keamanan", f"PIN seragam {pin} hanya untuk keperluan clustering/pelatihan. Setelah selesai, rotasi PIN "
                     "lewat menu Pengaturan website → Kredensial pegawai atau minta setiap pengguna mengganti PIN."),
    ]
    current = 3
    for title, body in steps:
        cell = guide.cell(row=current, column=1, value=title)
        cell.font = Font(bold=True, color="1F3864")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        body_cell = guide.cell(row=current, column=2, value=body)
        body_cell.alignment = Alignment(vertical="top", wrap_text=True)
        current += 1

    current += 1
    guide.cell(row=current, column=1, value="KODE PERAN & HAK AKSES").font = Font(bold=True, size=12, color="1F3864")
    current += 1
    for col, title in enumerate(["Label pada Excel", "Kode database / hak akses"], start=1):
        cell = guide.cell(row=current, column=col, value=title)
        cell.font = head_font
        cell.fill = header_fill
        cell.border = border
    current += 1
    for code, label in ROLES:
        guide.cell(row=current, column=1, value=label).border = border
        cell = guide.cell(row=current, column=2, value=f"{code} — {ROLE_ACCESS[code]}")
        cell.border = border
        cell.alignment = wrap
        current += 1

    current += 1
    guide.cell(row=current, column=1, value="DEPARTEMEN RESMI").font = Font(bold=True, size=12, color="1F3864")
    current += 1
    guide.cell(row=current, column=2,
               value="Admin; Bendahara; Perencanaan; Informasi dan Humas; Layanan Rehabilitasi Medis; "
                     "Layanan Rehabilitasi Sosial; Umum; Sarana & Prasarana; Clinical Supervisor").alignment = wrap

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


# --------------------------------------------------------------------------- #
# pin
# --------------------------------------------------------------------------- #
def apply_pin(store: Snapshot, pin: str, only_ids: Optional[Sequence[str]] = None,
              revoke_sessions: bool = True) -> Dict[str, Any]:
    if not (len(pin) == 6 and pin.isdigit()):
        raise SystemExit("PIN harus 6 digit angka.")

    targets = roster_staff(store)
    if only_ids:
        wanted = set(only_ids)
        targets = [row for row in targets if row.get("id") in wanted]

    fields = uniform_pin_fields(pin)
    changed = 0
    for row in targets:
        row.update(fields)
        if row.get("identity_is_placeholder") and not row.get("credential_source"):
            row["credential_source"] = "uniform-clustering"
        changed += 1

    if revoke_sessions:
        store.data["auth_sessions"] = []

    return {"pin": pin, "accounts": changed, "sessions_revoked": bool(revoke_sessions)}


# --------------------------------------------------------------------------- #
# apply (peran dari Excel)
# --------------------------------------------------------------------------- #
TRUE_MARKS = {"1", "01", "x", "X", "✓", "v", "V", "ya", "YA", "yes", "TRUE", "True", "true"}


def _is_marked(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return int(value) == 1
    return str(value).strip() in TRUE_MARKS


def read_template(path: Path) -> Tuple[List[Dict[str, Any]], Optional[Dict[str, Any]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    if "Isi Peran" not in wb.sheetnames:
        raise SystemExit("Sheet 'Isi Peran' tidak ditemukan pada berkas ini.")

    ws = wb["Isi Peran"]
    header_row = None
    for row in range(1, 15):
        if str(ws.cell(row=row, column=2).value or "").strip() == "ID Akun":
            header_row = row
            break
    if header_row is None:
        raise SystemExit("Baris header ('ID Akun') tidak ditemukan pada sheet 'Isi Peran'.")

    role_cols = {}
    for col in range(1, ws.max_column + 1):
        label = str(ws.cell(row=header_row, column=col).value or "").strip()
        if label in ROLE_CODE_BY_LABEL:
            role_cols[ROLE_CODE_BY_LABEL[label]] = col
    missing = [label for code, label in ROLES if code not in role_cols]
    if missing:
        raise SystemExit(f"Kolom peran tidak lengkap: {', '.join(missing)}")

    entries: List[Dict[str, Any]] = []
    blank_streak = 0
    for row in range(header_row + 1, ws.max_row + 1):
        marker = str(ws.cell(row=row, column=1).value or "")
        if marker.startswith("AKUN SUPER ADMIN TERPISAH"):
            break
        ident = str(ws.cell(row=row, column=2).value or "").strip()
        if not ident:
            blank_streak += 1
            if blank_streak >= 5:
                break
            continue
        blank_streak = 0
        marks = [code for code, col in role_cols.items() if _is_marked(ws.cell(row=row, column=col).value)]
        entries.append({
            "row": row,
            "id": ident,
            "nama": str(ws.cell(row=row, column=3).value or "").strip(),
            "nip": normalize_digits(ws.cell(row=row, column=4).value),
            "department": str(ws.cell(row=row, column=5).value or "").strip(),
            "marks": marks,
            "supervised": str(ws.cell(row=row, column=role_cols["staff"] + 1).value or "").strip(),
        })

    separate: Optional[Dict[str, Any]] = None
    for row in range(1, ws.max_row + 1):
        marker = str(ws.cell(row=row, column=1).value or "")
        if marker.startswith("AKUN SUPER ADMIN TERPISAH"):
            values = [ws.cell(row=row + 2, column=col).value for col in range(1, 5)]
            separate = {
                "name": str(values[0] or "").strip(),
                "nip": normalize_digits(values[1]),
                "pin": normalize_digits(values[2]) or DEFAULT_PIN,
                "note": str(values[3] or "").strip(),
                "row": row + 2,
            }
            break
    return entries, separate


def apply_roles(store: Snapshot, entries: Iterable[Dict[str, Any]], pin: str,
                strict: bool = False) -> Dict[str, Any]:
    staff_by_id = {str(row.get("id")): row for row in store.staff}
    seen_nip: Dict[str, str] = {}
    for row in store.staff:
        nip = normalize_digits(row.get("nip"))
        if nip:
            seen_nip[nip] = str(row.get("id"))

    problems: List[str] = []
    warnings: List[str] = []
    missing_identity: List[Tuple[str, str]] = []
    changed: List[Tuple[str, str, str]] = []
    fields = uniform_pin_fields(pin)

    for entry in entries:
        ident = entry["id"]
        row = staff_by_id.get(ident)
        if row is None:
            problems.append(f"Baris {entry['row']}: ID '{ident}' tidak ada di database.")
            continue
        if row.get("system_super_admin"):
            warnings.append(f"Baris {entry['row']}: '{ident}' adalah akun Super Admin terpisah; dilewati.")
            continue
        if len(entry["marks"]) > 1:
            labels = ", ".join(ROLE_LABEL_BY_CODE[code] for code in entry["marks"])
            problems.append(
                f"Baris {entry['row']} ({ident}): lebih dari satu peran ditandai ({labels}). "
                f"Hapus angka 1 pada kolom peran lain (pilih sel lalu tekan Delete), sisakan satu."
            )
            continue
        if not entry["marks"]:
            if strict:
                problems.append(f"Baris {entry['row']} ({ident}): belum ada peran yang ditandai.")
                continue
            warnings.append(f"Baris {entry['row']} ({ident}): peran kosong, peran lama dipertahankan.")
            role_code = str(row.get("role") or "staff")
        else:
            role_code = entry["marks"][0]

        nip = entry["nip"]
        if nip:
            if len(nip) not in (16, 18):
                problems.append(f"Baris {entry['row']} ({ident}): NIP/NIK '{nip}' bukan 16/18 digit.")
                continue
            owner = seen_nip.get(nip)
            if owner and owner != ident:
                problems.append(f"Baris {entry['row']} ({ident}): NIP/NIK '{nip}' sudah dipakai {owner}.")
                continue
        else:
            nip = normalize_digits(row.get("nip"))
            if not nip:
                missing_identity.append((ident, str(row.get("name") or "")))

        supervised = [
            item.strip() for item in entry["supervised"].replace(",", ";").split(";") if item.strip()
        ]
        if role_code not in SUPERVISOR_ROLES:
            supervised = []
        elif not supervised:
            supervised = [row.get("department")] if row.get("department") else []

        before = str(row.get("role") or "staff")
        row["role"] = role_code
        if nip:
            row["nip"] = nip
            row["id_type"] = id_type_for(nip)
            seen_nip[nip] = ident
        row["supervised_departments"] = supervised
        row.update(fields)
        if before != role_code:
            changed.append((ident, before, role_code))

    result = {
        "pin": pin,
        "problems": problems,
        "warnings": warnings,
        "missing_identity": missing_identity,
        "role_changes": changed,
        "processed": len([entry for entry in entries]),
    }
    return result


def upsert_separate_super_admin(store: Snapshot, info: Optional[Dict[str, Any]]) -> Dict[str, Any]:
    if not info or not info.get("nip"):
        return {"created": False, "reason": "NIP/NIK Super Admin terpisah kosong — dilewati."}
    nip = normalize_digits(info["nip"])
    if len(nip) not in (16, 18):
        return {"created": False, "reason": f"NIP/NIK Super Admin '{nip}' bukan 16/18 digit — dilewati."}
    for row in store.staff:
        if str(row.get("id")) != SEPARATE_SUPER_ADMIN_ID and normalize_digits(row.get("nip")) == nip:
            return {"created": False, "reason": f"NIP/NIK {nip} sudah dipakai akun {row.get('id')}."}

    name = info.get("name") or SEPARATE_SUPER_ADMIN_DEFAULT["name"]
    pin = info.get("pin") or DEFAULT_PIN
    existing = next(
        (row for row in store.staff
         if str(row.get("id")) == SEPARATE_SUPER_ADMIN_ID or row.get("system_super_admin")),
        None,
    )
    payload = {
        "id": SEPARATE_SUPER_ADMIN_ID,
        "name": name,
        "initials": "".join(part[0] for part in name.split()[:2]).upper() or "SA",
        "department": "Admin",
        "position": "Super Admin",
        "keterangan": "Akun sistem terpisah (di luar roster staf)",
        "active": True,
        "archived": False,
        "nip": nip,
        "id_type": id_type_for(nip),
        "role": "super_admin",
        "system_super_admin": True,
        "supervised_departments": [],
        "identity_is_placeholder": True,
        "credential_source": "clustering-setup",
    }
    payload.update(uniform_pin_fields(pin, source="clustering-setup"))
    if existing is None:
        store.staff.append(payload)
        return {"created": True, "id": SEPARATE_SUPER_ADMIN_ID, "pin": pin, "nip": nip}
    existing.update(payload)
    return {"created": False, "updated": True, "id": SEPARATE_SUPER_ADMIN_ID, "pin": pin, "nip": nip}


# --------------------------------------------------------------------------- #
# report
# --------------------------------------------------------------------------- #
def print_report(store: Snapshot) -> None:
    staff = roster_staff(store)
    separate = [row for row in store.staff if row.get("system_super_admin")]
    print(f"Total akun roster : {len(staff)}")
    role_counts = Counter(str(row.get("role") or "staff") for row in staff)
    print("\nSebaran peran (roster):")
    for code, label in ROLES:
        count = role_counts.get(code, 0)
        if code == "super_admin" and not count:
            continue
        print(f"  {label:<20} {count:>3}")
    print(f"  {'Super Admin terpisah':<20} {len(separate):>3}")

    print("\nMatriks departemen x peran:")
    departments = sorted({str(row.get("department") or "-") for row in staff})
    width = max(14, max(len(label) for _code, label in ROLES) + 3)
    header = f"  {'Departemen':<32}" + "".join(f"{label:>{width}}" for _code, label in ROLES)
    print(header)
    for department in departments:
        cells = []
        for code, _label in ROLES:
            count = sum(1 for row in staff
                        if str(row.get("department") or "-") == department
                        and str(row.get("role") or "staff") == code)
            cells.append(f"{count if count else '-':>{width}}")
        print(f"  {department:<32}" + "".join(cells))

    missing_nip = [row for row in staff if len(normalize_digits(row.get("nip"))) not in (16, 18)]
    belum_aktif = [row for row in staff if not row.get("is_activated")]
    print(f"\nAkun tanpa NIP/NIK valid (tidak bisa login): {len(missing_nip)}")
    for row in missing_nip[:10]:
        print(f"  - {row.get('id')}  {row.get('name')}")
    if len(missing_nip) > 10:
        print(f"  ... dan {len(missing_nip) - 10} akun lain")
    print(f"Akun belum aktivasi: {len(belum_aktif)}")
    if separate:
        for row in separate:
            print(f"Super Admin terpisah: {row.get('name')} | NIP/NIK {row.get('nip')} | id {row.get('id')}")


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Clustering akun LOKA-Kin")
    parser.add_argument("--data-file", type=Path, default=DATA_FILE,
                        help=f"snapshot database (default {DATA_FILE})")
    sub = parser.add_subparsers(dest="command", required=True)

    tpl = sub.add_parser("template", help="buat berkas Excel untuk mengisi peran")
    tpl.add_argument("--out", type=Path, default=DEFAULT_TEMPLATE, help=f"lokasi berkas (default {DEFAULT_TEMPLATE.name})")
    tpl.add_argument("--pin", default=DEFAULT_PIN, help=f"PIN yang ditulis pada panduan (default {DEFAULT_PIN})")

    pin_cmd = sub.add_parser("pin", help="seragamkan PIN semua akun roster")
    pin_cmd.add_argument("--pin", default=DEFAULT_PIN, help=f"PIN 6 digit (default {DEFAULT_PIN})")
    pin_cmd.add_argument("--keep-sessions", action="store_true", help="jangan cabut sesi login yang ada")
    pin_cmd.add_argument("--dry-run", action="store_true", help="tampilkan rencana tanpa menulis")

    app = sub.add_parser("apply", help="terapkan peran dari berkas Excel")
    app.add_argument("path", nargs="?", type=Path, default=DEFAULT_TEMPLATE, help="berkas Excel terisi")
    app.add_argument("--pin", default=DEFAULT_PIN, help=f"PIN seragam (default {DEFAULT_PIN})")
    app.add_argument("--strict", action="store_true", help="gagalkan bila ada baris tanpa peran")
    app.add_argument("--keep-sessions", action="store_true", help="jangan cabut sesi login yang ada")
    app.add_argument("--no-backup", action="store_true", help="lewati pencadangan snapshot")

    sub.add_parser("report", help="tampilkan matriks cluster dari database")

    args = parser.parse_args(argv)
    store = Snapshot(args.data_file)

    if args.command == "template":
        path = build_template(store, args.out, args.pin)
        print(f"Template dibuat: {path}")
        print(f"  staf roster: {len(roster_staff(store))} akun")
        return 0

    if args.command == "report":
        print_report(store)
        return 0

    if args.command == "pin":
        plan = apply_pin(store, args.pin, revoke_sessions=not args.keep_sessions)
        if args.dry_run:
            print(f"[dry-run] {plan['accounts']} akun akan memakai PIN {plan['pin']}")
            return 0
        backup = store.backup()
        store.save()
        print(f"PIN {plan['pin']} diterapkan pada {plan['accounts']} akun roster.")
        print(f"Sesi login dicabut: {'ya' if plan['sessions_revoked'] else 'tidak'}")
        print(f"Cadangan: {backup}")
        return 0

    path = Path(args.path)
    if not path.is_file():
        raise SystemExit(f"Berkas tidak ditemukan: {path}")
    entries, separate = read_template(path)
    result = apply_roles(store, entries, args.pin, strict=args.strict)
    if result["problems"]:
        print("Peran TIDAK diterapkan karena ada masalah pada berkas:")
        for problem in result["problems"]:
            print(f"  - {problem}")
        print("Perbaiki lalu jalankan ulang perintah apply.")
        return 1

    super_admin = upsert_separate_super_admin(store, separate)
    if not args.keep_sessions:
        store.data["auth_sessions"] = []
    backup = store.backup() if not args.no_backup else None
    store.save()

    print(f"Berhasil menerapkan {result['processed']} baris. PIN seragam: {result['pin']}")
    if result["role_changes"]:
        print("Perubahan peran:")
        for ident, before, after in result["role_changes"]:
            print(f"  - {ident}: {ROLE_LABEL_BY_CODE.get(before, before)} -> {ROLE_LABEL_BY_CODE.get(after, after)}")
    if result["warnings"]:
        print("Peringatan:")
        for warning in result["warnings"][:20]:
            print(f"  - {warning}")
        if len(result["warnings"]) > 20:
            print(f"  ... dan {len(result['warnings']) - 20} peringatan lain")
    missing = result.get("missing_identity") or []
    if missing:
        print(f"Akun tanpa NIP/NIK (peran disimpan, login belum bisa): {len(missing)}")
        for ident, name in missing[:5]:
            print(f"  - {ident}  {name}")
        if len(missing) > 5:
            print(f"  ... dan {len(missing) - 5} akun lain; lengkapi kolom NIP/NIK di sheet 'Isi Peran'.")
    if super_admin.get("created") or super_admin.get("updated"):
        action = "dibuat" if super_admin.get("created") else "diperbarui"
        print(f"Super Admin terpisah {action}: {super_admin['id']} | NIP/NIK {super_admin['nip']} | PIN {super_admin['pin']}")
    elif super_admin.get("reason"):
        print(f"Super Admin terpisah: {super_admin['reason']}")
    if backup:
        print(f"Cadangan: {backup}")
    print()
    print_report(store)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
