#!/usr/bin/env python3
"""Pratinjau LOKA-Kin satu-origin + utilitas data demo.

Sub-perintah:
  serve   Jalankan API FastAPI sekaligus menyajikan `frontend/build` pada
          origin yang sama (default http://0.0.0.0:8000). Satu origin membuat
          permintaan `/api` relatif dan cookie sesi selalu sejalan.
  proxy   Proxy ringan 3000 -> 8000 agar pratinjau dapat dibuka pada port itu
          juga tanpa menjalankan penulis database kedua.
  demo    Siapkan data demo:
            --roster    (jalankan saat server berhenti) semai roster resmi
                        lalu tetapkan NIP/role/PIN akun demo pada snapshot.
            --accounts  (jalankan saat server hidup) lengkapi kredensial seluruh
                        staf lewat alur resmi dan tulis daftar kredensial (.docx).

Cookie sesi diberi atribut SameSite=None; Secure; Partitioned (CHIPS) agar
login tetap bekerja ketika aplikasi dibuka di dalam iframe lintas-situs (mis.
panel preview), tanpa membagikan token sesi antarpengguna.

Konfigurasi lewat variabel lingkungan:
  AUTH_SETUP_SECRET   rahasia penyiapan Admin pertama (default: loka-preview-setup)
  PREVIEW_HOST/PORT   alamat `serve`       (default 0.0.0.0 / 8000)
  PROXY_PORT          port `proxy`         (default 3000)
  DEMO_ADMIN_NIP/PIN  akun Admin untuk `demo --accounts`
                      (default mengikuti admin demo bawaan 197804122006041005 / 246810)
"""

from __future__ import annotations

import argparse
import asyncio
import http.client
import json
import os
import random
import re
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
BACKEND = REPO / "backend"
FRONTEND_BUILD = REPO / "frontend" / "build"
WORK = REPO / ".preview"
CREDENTIALS = WORK / "kredensial.json"
DOCX_DEFAULT = Path.home() / "Kredensial-Staf-LOKA-Kin.docx"
XLSX_DEFAULT = Path.home() / "PIN-Staf-LOKA-Kin.xlsx"

sys.path.insert(0, str(BACKEND))

# Akun demo dengan identitas contoh yang ditandai non-produksi dan PIN tetap.
DEMO_ACCOUNTS = {
    "staff-9": {"nama": "Daniel, A.Md.Kep", "nip": "197804122006041005", "role": "super_admin",
                "department": "Admin", "supervised": [], "pin": "246810"},
    "staff-3": {"nama": "dr. Heni Purwanti", "nip": "197501012005012003", "role": "team_lead",
                "department": "Layanan Rehabilitasi Medis",
                "supervised": ["Layanan Rehabilitasi Medis"], "pin": "778899"},
    "staff-2": {"nama": "Nurma Fitria, S.IP, M.IKom", "nip": "198611232010012008", "role": "staff",
                "department": "Umum", "supervised": [], "pin": "135790"},
    "staff-1": {"nama": "Edwin, S.Sos", "nip": "1801010101010001", "role": "staff",
                "department": "Layanan Rehabilitasi Sosial", "supervised": [], "pin": "112233"},
}

TUGAS_DEMO = {
    "staff-2": [("Rekap absensi pegawai minggu ini", "5 hari kerja", "Tinggi"),
                ("Distribusi surat masuk ke unit terkait", "20 surat", "Sedang"),
                ("Perbarui arsip kepegawaian", "1 berkas", "Rendah")],
    "staff-1": [("Pendampingan konseling lanjutan klien", "6 sesi", "Tinggi"),
                ("Rekap catatan kegiatan harian unit", "1 laporan", "Sedang")],
}

PIN_DILARANG = {"123456", "111111", "222222", "333333", "444444", "555555",
                "666666", "777777", "888888", "999999", "000000"}
LABEL_PERAN = {"admin": "Admin", "super_admin": "Super Admin", "team_lead": "Ketua Tim", "clinical_supervisor": "Clinical Supervisor",
               "head": "Kepala", "staff": "Staf"}


# --------------------------------------------------------------------------- #
# serve / proxy
# --------------------------------------------------------------------------- #
def _bangun_app_preview():
    """Aplikasi LOKA-Kin + penyajian hasil build + penyesuaian cookie iframe."""
    import server  # noqa: WPS433 (impor aplikasi FastAPI proyek)
    from fastapi.responses import FileResponse, JSONResponse
    from fastapi.staticfiles import StaticFiles

    WORK.mkdir(parents=True, exist_ok=True)
    log_path = WORK / "requests.log"
    app = server.app

    class PreviewSupport:
        def __init__(self, app, **options):
            self.app = app

        async def __call__(self, scope, receive, send):
            if scope["type"] != "http":
                return await self.app(scope, receive, send)
            headers = {k.decode("latin-1").lower(): v.decode("latin-1") for k, v in scope.get("headers", [])}
            try:
                with log_path.open("a", encoding="utf-8") as fh:
                    fh.write(json.dumps({
                        "t": time.strftime("%H:%M:%S"), "method": scope.get("method"),
                        "path": scope.get("path"), "host": headers.get("host"),
                        "xf_host": headers.get("x-forwarded-host"),
                        "ua": (headers.get("user-agent") or "")[:70],
                        # Presence only: never log credentials, cookies or tokens.
                        "bearer_present": headers.get("authorization", "").lower().startswith("bearer "),
                        "session_cookie_present": "loka_kin_session=" in headers.get("cookie", ""),
                    }, ensure_ascii=False) + "\n")
            except Exception:
                pass

            async def send_wrapper(message):
                if message["type"] == "http.response.start":
                    baru = []
                    for key, value in message.get("headers", []):
                        if key.lower() == b"set-cookie":
                            teks = value.decode("latin-1")
                            teks = re.sub(r"(?i)samesite\s*=\s*lax", "SameSite=None; Secure", teks)
                            if "samesite=" not in teks.lower():
                                teks = f"{teks}; SameSite=None; Secure"
                            if "partitioned" not in teks.lower():
                                teks = f"{teks}; Partitioned"
                            value = teks.encode("latin-1")
                        baru.append((key, value))
                    if not scope.get("path", "").startswith("/api"):
                        baru = [(key, value) for key, value in baru if key.lower() != b"cache-control"]
                        baru.append((b"cache-control", b"no-store"))
                    message["headers"] = baru
                await send(message)

            await self.app(scope, receive, send_wrapper)

    app.add_middleware(PreviewSupport)

    if (FRONTEND_BUILD / "static").is_dir():
        app.mount("/static", StaticFiles(directory=str(FRONTEND_BUILD / "static")), name="spa-static")

    @app.get("/", include_in_schema=False)
    @app.get("/{full_path:path}", include_in_schema=False)
    async def spa(full_path: str = ""):
        if full_path.startswith("api"):
            return JSONResponse({"detail": "Not Found"}, status_code=404)
        index = FRONTEND_BUILD / "index.html"
        kandidat = FRONTEND_BUILD / full_path
        if full_path and kandidat.is_file() and ".." not in full_path:
            return FileResponse(kandidat)
        if index.is_file():
            return FileResponse(index, headers={"Cache-Control": "no-store"})
        return JSONResponse({"detail": "Frontend belum dibangun. Jalankan: cd frontend && yarn build"}, status_code=503)

    return app


def serve() -> None:
    import uvicorn

    if not (FRONTEND_BUILD / "index.html").is_file():
        print(f"[peringatan] {FRONTEND_BUILD} belum ada. Jalankan: cd frontend && yarn build")
    host = os.environ.get("PREVIEW_HOST", "0.0.0.0")
    port = int(os.environ.get("PREVIEW_PORT", "8000"))
    print(f"LOKA-Kin satu-origin http://{host}:{port} (API + dashboard)")
    uvicorn.run(_bangun_app_preview(), host=host, port=port, log_level="info")


def proxy() -> None:
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

    upstream_host = os.environ.get("PREVIEW_UPSTREAM_HOST", "127.0.0.1")
    upstream_port = int(os.environ.get("PREVIEW_PORT", "8000"))
    port = int(os.environ.get("PROXY_PORT", "3000"))
    hop = {"connection", "keep-alive", "proxy-authenticate", "proxy-authorization",
           "te", "trailers", "transfer-encoding", "upgrade"}

    class Handler(BaseHTTPRequestHandler):
        protocol_version = "HTTP/1.1"
        server_version = "LokaKinPreviewProxy"

        def log_message(self, *args):  # senyapkan log per-permintaan
            pass

        def _forward(self):
            length = int(self.headers.get("Content-Length") or 0)
            body = self.rfile.read(length) if length else None
            headers = {k: v for k, v in self.headers.items() if k.lower() not in hop}
            headers["Host"] = f"{upstream_host}:{upstream_port}"
            headers["X-Forwarded-Proto"] = "https"
            headers["X-Forwarded-Host"] = self.headers.get("Host", "")
            try:
                conn = http.client.HTTPConnection(upstream_host, upstream_port, timeout=60)
                conn.request(self.command, self.path, body=body, headers=headers)
                response = conn.getresponse()
                payload = response.read()
            except Exception as exc:
                pesan = f"Upstream tidak dapat dihubungi: {exc}".encode()
                self.send_response(502)
                self.send_header("Content-Type", "text/plain; charset=utf-8")
                self.send_header("Content-Length", str(len(pesan)))
                self.end_headers()
                self.wfile.write(pesan)
                return
            self.send_response(response.status)
            for key, value in response.getheaders():
                if key.lower() in hop or key.lower() == "content-length":
                    continue
                self.send_header(key, value)
            self.send_header("Content-Length", str(len(payload)))
            self.end_headers()
            if self.command != "HEAD":
                self.wfile.write(payload)

        do_GET = do_POST = do_PUT = do_PATCH = do_DELETE = do_HEAD = do_OPTIONS = _forward

    print(f"LOKA-Kin proxy http://0.0.0.0:{port} -> {upstream_host}:{upstream_port}")
    ThreadingHTTPServer(("0.0.0.0", port), Handler).serve_forever()


# --------------------------------------------------------------------------- #
# demo: roster
# --------------------------------------------------------------------------- #
async def demo_roster() -> None:
    from auth_utils import hash_credential
    import server as backend_server

    # PENTING: gunakan instance database milik modul server. Membuat instance
    # baru akan menghasilkan penyimpanan mongomock in-memory yang terpisah,
    # sehingga hasil penyemaian server.seed_data tidak akan terlihat di sini.
    db = backend_server.db
    print(f"[roster] mode={backend_server.DB_MODE}")
    dipulihkan = await db.restore()
    print(f"[roster] dokumen dipulihkan: {dipulihkan}")
    await backend_server.seed_data(force=True)  # semai + sinkronkan roster resmi
    print("[roster] staf tersemai:", await db.staff.count_documents({}))

    for staff_id, info in DEMO_ACCOUNTS.items():
        doc = await db.staff.find_one({"id": staff_id}, {"_id": 0})
        if not doc:
            print(f"[roster] {staff_id} tidak ditemukan")
            continue
        await db.staff.update_one({"id": staff_id}, {"$set": {
            "name": doc["name"], "nip": info["nip"], "role": info["role"],
            "department": info["department"], "supervised_departments": info["supervised"],
            "pin_hash": hash_credential(info["pin"]), "is_activated": True,
            "activation_code_hash": "", "activation_code_expires_at": "",
            "activation_code_attempts": 0, "activation_code_locked_until": "",
            "login_failed_attempts": 0, "login_locked_until": "", "seed_key": "",
            "identity_is_placeholder": True, "credential_source": "demo-preview",
        }})
        print(f"[roster] {staff_id}: {doc['name'][:28]:30} {info['role']:9} PIN {info['pin']}")
    print("[roster] staf:", await db.staff.count_documents({}), "| tugas:", await db.tasks.count_documents({}))


# --------------------------------------------------------------------------- #
# demo: accounts (lewat API resmi)
# --------------------------------------------------------------------------- #
API = "http://127.0.0.1:8000"


def _panggil(method: str, path: str, payload=None, token: str = ""):
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Cookie"] = f"loka_kin_session={token}"
    req = urllib.request.Request(API + path, method=method, headers=headers,
                                 data=json.dumps(payload).encode() if payload is not None else None)
    try:
        with urllib.request.urlopen(req) as r:
            body = r.read().decode()
            return r.status, (json.loads(body) if body.strip() else None), r.headers
    except urllib.error.HTTPError as e:
        body = e.read().decode()
        try:
            body = json.loads(body)
        except Exception:
            pass
        return e.code, body, e.headers


class SesiAdmin:
    """Sesi Admin yang menyegarkan token sendiri (sesi dapat dicabut saat kode aktivasi terbit).

    Penyegaran dibatasi: minimal 5 detik antar percobaan dan maksimal 3 kali
    beruntun, supaya kegagalan otorisasi tidak memicu penguncian login (429).
    """

    def __init__(self) -> None:
        self.nip = os.environ.get("DEMO_ADMIN_NIP") or DEMO_ACCOUNTS["staff-9"]["nip"]
        self.pin = os.environ.get("DEMO_ADMIN_PIN") or DEMO_ACCOUNTS["staff-9"]["pin"]
        self.token = ""
        self._segar_terakhir = 0.0
        self._gagal_segar = 0

    def masuk(self) -> bool:
        status, _, headers = _panggil("POST", "/api/auth/login", {"nip": self.nip, "pin": self.pin})
        if status != 200:
            return False
        self.token = headers.get("set-cookie", "").split("loka_kin_session=")[1].split(";")[0]
        return True

    def __call__(self, method: str, path: str, payload=None):
        hasil = _panggil(method, path, payload, token=self.token)
        if hasil[0] in (401, 403) and time.time() - self._segar_terakhir > 5 and self._gagal_segar < 3:
            self._segar_terakhir = time.time()
            if self.masuk():
                self._gagal_segar = 0
                hasil = _panggil(method, path, payload, token=self.token)
            else:
                self._gagal_segar += 1
        return hasil


def demo_accounts(docx: Path | None, xlsx: Path | None = None) -> int:
    admin = SesiAdmin()
    if not admin.masuk():
        secret = os.environ.get("AUTH_SETUP_SECRET", "loka-preview-setup")
        status, body, _ = _panggil("POST", "/api/auth/bootstrap-admin", {"setup_secret": secret, "nip": admin.nip})
        print("[accounts] bootstrap Admin pertama:", status)
        if status == 200 and body.get("activation_code"):
            aktif = _panggil("POST", "/api/auth/activate", {
                "nip": admin.nip, "activation_code": body["activation_code"],
                "pin": admin.pin, "pin_confirmation": admin.pin,
            })[0]
            print("[accounts] aktivasi Admin:", aktif)
        if not admin.masuk():
            print("[accounts] gagal masuk sebagai Admin — pastikan server hidup dan PIN benar")
            return 1

    _, daftar, _ = admin("GET", "/api/staff")
    daftar.sort(key=lambda s: int(s["id"].split("-")[-1]) if s["id"].split("-")[-1].isdigit() else 10**6)
    print(f"[accounts] staf pada roster: {len(daftar)}")

    lama = {sid: {"id": sid, "nip": info["nip"], "pin": info["pin"], "nama": info["nama"],
                  "sumber": "akun demo (PIN tetap)"}
            for sid, info in DEMO_ACCOUNTS.items()}
    if CREDENTIALS.exists():
        lama.update(json.loads(CREDENTIALS.read_text()))
    dipakai = {v["pin"] for v in lama.values() if v.get("pin")}
    hasil, gagal, baru, rusak = [], [], 0, []
    random.seed()

    for indeks, staf in enumerate(daftar, start=1):
        sid = staf["id"]
        simpan = lama.get(sid)

        # Akun demo (termasuk Admin yang sedang dipakai) TIDAK boleh diterbitkan
        # ulang: menerbitkan kode aktivasi untuk Admin akan mencabut sesi ini dan
        # menimpa PIN Admin. Perbaikannya lewat `demo --roster` saat server mati.
        if sid in DEMO_ACCOUNTS:
            if _panggil("POST", "/api/auth/login", {"nip": simpan["nip"], "pin": simpan["pin"]})[0] != 200:
                rusak.append(sid)
            hasil.append({**simpan, "nama": staf["name"], "department": staf["department"],
                          "role": staf["role"]})
            continue

        if simpan and _panggil("POST", "/api/auth/login", {"nip": simpan["nip"], "pin": simpan["pin"]})[0] == 200:
            hasil.append({**simpan, "nama": staf["name"], "department": staf["department"],
                          "role": staf["role"], "sumber": "sudah ada"})
            continue

        nip = staf.get("nip") or (simpan or {}).get("nip")
        if not nip:
            gagal.append((staf["name"], "NIP/NIK resmi belum tersedia; impor roster resmi sebelum aktivasi"))
            continue

        status, body, _ = admin("POST", f"/api/auth/staff/{sid}/activation-code")
        if status in (401, 403, 429):
            print(f"[accounts] otorisasi Admin bermasalah ({status}) — proses dihentikan. "
                  "Hentikan server, jalankan `preview.py demo --roster`, lalu ulangi.")
            break
        if status != 200:
            gagal.append((staf["name"], f"kode aktivasi gagal ({status})"))
            continue

        pin = (simpan or {}).get("pin") or ""
        if not pin or pin in dipakai:
            while True:
                pin = f"{random.randint(0, 999999):06d}"
                if pin not in dipakai and pin not in PIN_DILARANG:
                    break
        dipakai.add(pin)

        status, body, _ = _panggil("POST", "/api/auth/activate", {
            "nip": nip, "activation_code": body["activation_code"], "pin": pin, "pin_confirmation": pin})
        if status != 200:
            gagal.append((staf["name"], f"aktivasi gagal ({status})"))
            continue
        if _panggil("POST", "/api/auth/login", {"nip": nip, "pin": pin})[0] != 200:
            gagal.append((staf["name"], "login uji gagal"))
            continue

        baru += 1
        hasil.append({"id": sid, "nip": nip, "pin": pin, "nama": staf["name"],
                      "department": staf["department"], "role": staf["role"], "sumber": "baru"})

    WORK.mkdir(parents=True, exist_ok=True)
    CREDENTIALS.write_text(json.dumps({h["id"]: h for h in hasil}, ensure_ascii=False, indent=2))
    print(f"[accounts] kredensial siap: {len(hasil)} (baru {baru}) | gagal: {len(gagal)}")
    if rusak:
        print(f"[accounts] PERHATIAN: kredensial akun demo bermasalah: {', '.join(rusak)}")
        print("[accounts] jalankan ulang: hentikan server lalu `preview.py demo --roster`")
    for nama, alasan in gagal:
        print(f"  GAGAL: {nama} — {alasan}")

    # laporan contoh agar dashboard tidak kosong
    for staff_id, daftar_tugas in TUGAS_DEMO.items():
        for judul, target, prioritas in daftar_tugas:
            admin("POST", "/api/tasks", {"title": judul, "staff_id": staff_id, "status": "todo",
                                        "target": target, "priority": prioritas, "due_date": "Hari ini",
                                        "notes": "Laporan contoh untuk pratinjau", "proof_link": "",
                                        "photo_data": "", "photo_name": ""})

    if docx:
        tulis_docx(hasil, docx)
    if xlsx:
        tulis_xlsx(hasil, xlsx)
    return 0 if not gagal else 2


# --------------------------------------------------------------------------- #
# dokumen Word
# --------------------------------------------------------------------------- #
def tulis_docx(hasil: list[dict], tujuan: Path) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from datetime import datetime
    from zoneinfo import ZoneInfo

    def warna(cell, hex_warna):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), hex_warna)
        cell._tc.get_or_add_tcPr().append(shd)

    def header_ulang(row):
        tr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        tr.append(el)

    baris = sorted(hasil, key=lambda r: (r.get("role") != "admin", r.get("role") != "team_lead",
                                        int(r["id"].split("-")[-1]) if r["id"].split("-")[-1].isdigit() else 10**6))
    doc = Document()
    for sisi in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(doc.sections[0], sisi, Cm(1.6))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    judul = doc.add_paragraph()
    judul.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = judul.add_run("LOKA-Kin — Daftar Username & PIN Staf")
    run.bold = True
    run.font.size = Pt(17)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Loka Rehabilitasi Narkotika Kalianda\n"
                      f"Total: {len(baris)} staf  •  Dibuat: "
                      f"{datetime.now(ZoneInfo('Asia/Jakarta')).strftime('%d %B %Y, %H:%M')} WIB\n"
                      "Lingkungan: pratinjau/demo (bukan data produksi)")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x44, 0x55, 0x66)

    kotak = doc.add_table(rows=1, cols=1)
    kotak.style = "Table Grid"
    sel = kotak.rows[0].cells[0]
    warna(sel, "FFF4F4")
    run = sel.paragraphs[0].add_run("RAHASIA — JANGAN DISEBARLUASKAN")
    run.bold = True
    run.font.color.rgb = RGBColor(0xB0, 0x1B, 0x1B)
    for teks in (
        "PIN hanya disimpan sebagai hash (PBKDF2) di server sehingga tidak dapat dibaca ulang. PIN pada daftar ini dibuatkan Admin untuk keperluan uji coba; setiap staf sebaiknya menggantinya lewat alur reset (Admin menerbitkan kode aktivasi baru).",
        "Kolom Username memuat NIP/NIK. Staf yang belum memiliki NIP/NIK pada daftar resmi memakai nomor contoh berpola 18xxxxxxxxxxxxxx yang WAJIB diganti dengan NIP/NIK resmi sebelum dipakai di produksi.",
        "Serahkan PIN hanya kepada staf yang bersangkutan, dan jangan letakkan berkas ini pada repositori kode, grup percakapan, atau folder bersama.",
    ):
        par = sel.add_paragraph()
        r = par.add_run("• " + teks)
        r.font.size = Pt(8.5)

    doc.add_paragraph()
    pimpinan = [r for r in baris if r.get("role") in {"admin", "team_lead", "clinical_supervisor", "head"}]
    if pimpinan:
        h = doc.add_paragraph()
        hb = h.add_run("Akun pimpinan (monitoring seluruh staf)")
        hb.bold = True
        hb.font.size = Pt(11)
        tp = doc.add_table(rows=1, cols=5)
        tp.style = "Table Grid"
        for sel_k, teks in zip(tp.rows[0].cells, ("No", "Nama", "Peran", "Departemen", "Username (NIP/NIK)")):
            warna(sel_k, "E8F0FF")
            r = sel_k.paragraphs[0].add_run(teks)
            r.bold = True
            r.font.size = Pt(9)
        header_ulang(tp.rows[0])
        for no, item in enumerate(pimpinan, start=1):
            for cell, teks in zip(tp.add_row().cells, (str(no), item["nama"],
                                                       LABEL_PERAN.get(item["role"], item["role"]),
                                                       item.get("department", "-"), item["nip"])):
                cell.paragraphs[0].add_run(teks).font.size = Pt(9)
        doc.add_paragraph()

    h = doc.add_paragraph()
    hb = h.add_run(f"Daftar lengkap {len(baris)} staf")
    hb.bold = True
    hb.font.size = Pt(11)
    tabel = doc.add_table(rows=1, cols=6)
    tabel.style = "Table Grid"
    tabel.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, lebar in enumerate((Cm(1.0), Cm(5.2), Cm(3.4), Cm(2.2), Cm(3.6), Cm(1.6))):
        tabel.columns[i].width = lebar
    for sel_k, teks in zip(tabel.rows[0].cells, ("No", "Nama Staf", "Departemen", "Peran",
                                                 "Username (NIP/NIK)", "PIN")):
        warna(sel_k, "E8F0FF")
        r = sel_k.paragraphs[0].add_run(teks)
        r.bold = True
        r.font.size = Pt(9)
    header_ulang(tabel.rows[0])
    for no, item in enumerate(baris, start=1):
        sel = tabel.add_row().cells
        for cell, teks in zip(sel, (str(no), item["nama"], item.get("department", "-"),
                                    LABEL_PERAN.get(item.get("role", "staff"), item.get("role", "staff")),
                                    item["nip"], item["pin"])):
            run = cell.paragraphs[0].add_run(teks)
            run.font.size = Pt(8.5)
            if teks == item["pin"]:
                run.bold = True
        if no % 2 == 0:
            for cell in sel:
                warna(cell, "F7F9FC")

    doc.add_paragraph()
    h = doc.add_paragraph()
    hb = h.add_run("Cara masuk aplikasi")
    hb.bold = True
    hb.font.size = Pt(11)
    for nomor, teks in enumerate([
        "Buka aplikasi LOKA-Kin, lalu klik tombol Masuk pada halaman depan.",
        "Isi NIP/NIK (kolom Username) dan PIN sesuai baris nama Anda, lalu klik Masuk.",
        "Setelah berhasil, klik tombol Dashboard pada header untuk membuka halaman kinerja harian.",
        "Staf hanya melihat laporan miliknya sendiri; Admin, Ketua Tim, Clinical Supervisor, dan Kepala dapat memantau seluruh staf.",
        "Lupa PIN atau PIN bocor: minta Admin menerbitkan kode aktivasi baru (12 digit, 24 jam, sekali pakai), lalu pilih Aktivasi akun / lupa PIN dan buat PIN sendiri.",
    ], start=1):
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(3)
        par.add_run(f"{nomor}. {teks}").font.size = Pt(9)

    tujuan.parent.mkdir(parents=True, exist_ok=True)
    doc.save(tujuan)
    print(f"[accounts] dokumen Word: {tujuan} ({len(baris)} baris)")



# --------------------------------------------------------------------------- #
# berkas Excel
# --------------------------------------------------------------------------- #
def tulis_xlsx(hasil: list[dict], tujuan: Path) -> None:
    """Tulis daftar kredensial ke berkas Excel (.xlsx).

    NIP/NIK dan PIN ditulis sebagai TEKS (number_format '@') supaya Excel tidak
    mengubah nomor panjang menjadi notasi ilmiah atau memangkas angka nol di depan.
    """
    from datetime import datetime
    from zoneinfo import ZoneInfo

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    tipis = Side(style="thin", color="C9D3E0")
    bingkai = Border(left=tipis, right=tipis, top=tipis, bottom=tipis)
    kepala_fill = PatternFill("solid", fgColor="E8F0FF")
    baris_fill = PatternFill("solid", fgColor="F7F9FC")
    peringatan_fill = PatternFill("solid", fgColor="FFF4F4")
    tengah = Alignment(horizontal="center", vertical="center")
    kiri = Alignment(horizontal="left", vertical="center")
    bungkus = Alignment(horizontal="left", vertical="top", wrap_text=True)

    baris = sorted(hasil, key=lambda r: (r.get("role") != "admin", r.get("role") != "team_lead",
                                        int(r["id"].split("-")[-1]) if r["id"].split("-")[-1].isdigit() else 10**6))
    wb = Workbook()

    # ---- lembar 1: daftar kredensial -------------------------------------
    ws = wb.active
    ws.title = "Kredensial Staf"
    kolom = ["No", "Nama Staf", "Departemen", "Peran", "Username (NIP/NIK)", "PIN"]
    lebar = [6, 34, 26, 18, 24, 12]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(kolom))
    sel = ws.cell(row=1, column=1, value="LOKA-Kin — Daftar Username & PIN Staf")
    sel.font = Font(bold=True, size=14)
    sel.alignment = tengah
    ws.row_dimensions[1].height = 22

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(kolom))
    sel = ws.cell(row=2, column=1, value=(
        "Loka Rehabilitasi Narkotika Kalianda  •  Total: "
        f"{len(baris)} staf  •  Dibuat: {datetime.now(ZoneInfo('Asia/Jakarta')).strftime('%d %B %Y, %H:%M')} WIB"
        "  •  Lingkungan: pratinjau/demo (bukan data produksi)"))
    sel.font = Font(size=9, color="445566")
    sel.alignment = tengah

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(kolom))
    sel = ws.cell(row=3, column=1, value=(
        "RAHASIA — JANGAN DISEBARLUASKAN.  PIN hanya tersimpan sebagai hash (PBKDF2) di server sehingga "
        "tidak dapat dibaca ulang; PIN di daftar ini dibuatkan Admin untuk uji coba.  Staf yang belum punya "
        "NIP/NIK resmi memakai nomor contoh berpola 18xxxxxxxxxxxxxx yang WAJIB diganti sebelum produksi.  "
        "Serahkan PIN hanya kepada staf yang bersangkutan."))
    sel.font = Font(bold=True, size=9, color="B01B1B")
    sel.alignment = bungkus
    sel.fill = peringatan_fill
    for kolom_ke in range(1, len(kolom) + 1):
        ws.cell(row=3, column=kolom_ke).fill = peringatan_fill
    ws.row_dimensions[3].height = 58

    baris_awal = 5
    for indeks, judul in enumerate(kolom, start=1):
        sel = ws.cell(row=baris_awal, column=indeks, value=judul)
        sel.font = Font(bold=True, size=10)
        sel.fill = kepala_fill
        sel.alignment = tengah
        sel.border = bingkai
    ws.row_dimensions[baris_awal].height = 18

    for no, item in enumerate(baris, start=1):
        r = baris_awal + no
        nilai = [no, item["nama"], item.get("department", "-"),
                 LABEL_PERAN.get(item.get("role", "staff"), item.get("role", "staff")),
                 item["nip"], item["pin"]]
        for indeks, teks in enumerate(nilai, start=1):
            sel = ws.cell(row=r, column=indeks, value=teks)
            sel.border = bingkai
            sel.alignment = kiri if indeks in (2, 3, 4) else tengah
            if indeks in (5, 6):          # NIP/NIK & PIN wajib teks
                sel.number_format = "@"
            if indeks == 6:
                sel.font = Font(bold=True, size=10)
            if no % 2 == 0:
                sel.fill = baris_fill

    for indeks, lebar_kolom in enumerate(lebar, start=1):
        ws.column_dimensions[get_column_letter(indeks)].width = lebar_kolom
    ws.freeze_panes = ws.cell(row=baris_awal + 1, column=1)
    ws.auto_filter.ref = f"A{baris_awal}:{get_column_letter(len(kolom))}{baris_awal + len(baris)}"

    # ---- lembar 2: panduan ----------------------------------------------
    ps = wb.create_sheet("Cara Masuk")
    ps.column_dimensions["A"].width = 4
    ps.column_dimensions["B"].width = 110
    ps["B2"] = "Cara masuk aplikasi LOKA-Kin"
    ps["B2"].font = Font(bold=True, size=13)
    langkah = [
        "Buka aplikasi LOKA-Kin, lalu klik tombol Masuk pada halaman depan.",
        "Isi NIP/NIK (kolom Username) dan PIN sesuai baris nama Anda, lalu klik Masuk.",
        "Setelah berhasil, klik tombol Dashboard pada header untuk membuka halaman kinerja harian.",
        "Staf hanya melihat laporan miliknya sendiri; Admin, Ketua Tim, Clinical Supervisor, dan Kepala dapat memantau seluruh staf.",
        "Ekspor laporan: role Staf langsung membuat dokumen laporannya sendiri; pimpinan dapat memilih seluruh staf atau satu staf.",
        "Lupa PIN atau PIN bocor: minta Admin menerbitkan kode aktivasi baru (12 digit, berlaku 24 jam, sekali pakai), lalu pilih Aktivasi akun / lupa PIN dan buat PIN sendiri.",
        "Jangan letakkan berkas ini pada repositori kode, grup percakapan, atau folder bersama.",
    ]
    for nomor, teks in enumerate(langkah, start=1):
        ps.cell(row=3 + nomor, column=1, value=nomor).font = Font(bold=True, size=10)
        sel = ps.cell(row=3 + nomor, column=2, value=teks)
        sel.alignment = bungkus
        sel.font = Font(size=10)
        ps.row_dimensions[3 + nomor].height = 28

    tujuan.parent.mkdir(parents=True, exist_ok=True)
    wb.save(tujuan)
    print(f"[ekspor] berkas Excel: {tujuan} ({len(baris)} baris)")


def ekspor(kredensial: Path, docx: Path | None, xlsx: Path | None) -> int:
    """Tulis ulang berkas Word/Excel dari kredensial yang sudah terekam."""
    if not kredensial.exists():
        print(f"[ekspor] {kredensial} belum ada — jalankan `demo --accounts` lebih dahulu")
        return 1
    hasil = list(json.loads(kredensial.read_text()).values())
    if docx:
        tulis_docx(hasil, docx)
    if xlsx:
        tulis_xlsx(hasil, xlsx)
    return 0


# --------------------------------------------------------------------------- #
def main() -> int:
    parser = argparse.ArgumentParser(description="Pratinjau LOKA-Kin & utilitas data demo")
    sub = parser.add_subparsers(dest="perintah", required=True)
    sub.add_parser("serve", help="jalankan API + dashboard satu-origin")
    sub.add_parser("proxy", help="proxy 3000 -> 8000")
    demo = sub.add_parser("demo", help="siapkan data demo")
    aksi = demo.add_mutually_exclusive_group(required=True)
    aksi.add_argument("--roster", action="store_true", help="semai roster + akun demo (server berhenti)")
    aksi.add_argument("--accounts", action="store_true", help="lengkapi kredensial seluruh staf (server hidup)")
    demo.add_argument("--docx", type=Path, default=DOCX_DEFAULT, help=f"lokasi dokumen Word (default {DOCX_DEFAULT})")
    demo.add_argument("--xlsx", type=Path, default=XLSX_DEFAULT, help=f"lokasi berkas Excel (default {XLSX_DEFAULT})")
    exp = sub.add_parser("ekspor", help="tulis ulang Word/Excel dari kredensial yang terekam")
    exp.add_argument("--kredensial", type=Path, default=CREDENTIALS, help="berkas rekaman kredensial")
    exp.add_argument("--docx", type=Path, default=None, help="tulis dokumen Word ke lokasi ini")
    exp.add_argument("--xlsx", type=Path, default=XLSX_DEFAULT, help=f"tulis berkas Excel (default {XLSX_DEFAULT})")

    args = parser.parse_args()
    if args.perintah == "serve":
        serve()
        return 0
    if args.perintah == "proxy":
        proxy()
        return 0
    if args.perintah == "ekspor":
        return ekspor(args.kredensial, args.docx, args.xlsx)
    if args.roster:
        asyncio.run(demo_roster())
        return 0
    return demo_accounts(args.docx, args.xlsx)


if __name__ == "__main__":
    raise SystemExit(main())
